import csv
import io
import re
from datetime import datetime

from django.db import transaction
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from university.document_design import (ReportDocTemplate, document_styles, document_fonts,
    PageNumberCanvas, letterhead, get_branding, finish_worksheet)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable, Image as RLImage, Paragraph, SimpleDocTemplate, Table, TableStyle
)

from accounts.models import FacultyProfile
from university.models import AcademicTerm, ClassSchedule, Course, ExamRoom
from university.student_io import NumberedCanvas


DAY_LABELS = dict(ClassSchedule.Day.choices)
DAY_ORDER = [c for c, _ in ClassSchedule.Day.choices]

TIMETABLE_IMPORT_COLUMNS = [
    ("term_name", "Academic Year / Term", True),
    ("day", "Day", True),
    ("start_time", "Start Time", True),
    ("end_time", "End Time", True),
    ("course_code", "Course Code", True),
    ("room_name", "Venue / Room", True),
    ("session_type", "Session Type", False),
    ("course_name", "Course Name", False),
    ("department_code", "Department", False),
    ("programme_code", "Programme", False),
    ("faculty", "Faculty / Lecturer", False),
]

SAMPLE_TIMETABLE_ROW = {
    "term_name": "Fall 2026",
    "day": "Monday",
    "start_time": "09:00",
    "end_time": "10:30",
    "course_code": "CS101",
    "room_name": "Main Hall",
    "session_type": "Lecture",
    "course_name": "Introduction to Programming",
    "department_code": "CSE",
    "programme_code": "BT-CSE",
    "faculty": "FAC1000",
}


# ==============================================================================
# 1. EXPORT FUNCTIONS (CSV, EXCEL, PDF)
# ==============================================================================

def export_timetable_csv(queryset):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    headers = ["Academic Year / Term", "Day", "Start Time", "End Time", "Course Code",
              "Course Name", "Department", "Programme", "Faculty / Lecturer",
              "Venue / Room", "Session Type", "Status"]
    writer.writerow(headers)

    for s in queryset:
        writer.writerow([
            s.term.name if s.term else "—",
            s.get_day_display(),
            s.start_time.strftime("%H:%M"),
            s.end_time.strftime("%H:%M"),
            s.course.code,
            s.course.title,
            s.course.department.code if s.course.department else "—",
            s.course.program.code if s.course.program else "—",
            s.course.faculty.user.display_name if s.course.faculty else "TBA",
            s.room.name,
            s.get_session_type_display(),
            s.get_status_display(),
        ])

    return buffer.getvalue().encode("utf-8-sig")


def export_timetable_excel(queryset, site_name="University Management System"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.views.sheetView[0].showGridLines = True

    primary_color = "6C5CE7"
    header_fill_color = "4834D4"
    zebra_color = "EEF2FF"
    border_color = "CBD5E1"

    font_title = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=10, italic=True, color="E0E7FF")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10, color="1E293B")
    font_bold_data = Font(name="Arial", size=10, bold=True, color="1E293B")

    thin_border = Border(
        left=Side(style="thin", color=border_color), right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color), bottom=Side(style="thin", color=border_color))

    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"  {site_name.upper()} — CLASS TIMETABLE"
    title_cell.font = font_title
    title_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:L2")
    sub_cell = ws["A2"]
    sub_cell.value = f"  Generated: {timezone.now().strftime('%d %b %Y, %H:%M')} | Total Sessions: {queryset.count()}"
    sub_cell.font = font_sub
    sub_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    sub_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    headers = ["Term", "Day", "Start", "End", "Course Code", "Course Name",
              "Department", "Programme", "Faculty", "Venue", "Type", "Status"]
    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_num)
        cell.value = header
        cell.font = font_header
        cell.fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if header in ("Day", "Start", "End", "Status", "Type") else "left",
                                   vertical="center")
        cell.border = thin_border

    for row_idx, s in enumerate(queryset, start=5):
        ws.row_dimensions[row_idx].height = 22
        fill_color = zebra_color if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        values = [
            (s.term.name if s.term else "—", font_data, "left"),
            (s.get_day_display(), font_data, "center"),
            (s.start_time.strftime("%H:%M"), font_data, "center"),
            (s.end_time.strftime("%H:%M"), font_data, "center"),
            (s.course.code, font_bold_data, "left"),
            (s.course.title, font_data, "left"),
            (s.course.department.code if s.course.department else "—", font_data, "center"),
            (s.course.program.code if s.course.program else "—", font_data, "center"),
            (s.course.faculty.user.display_name if s.course.faculty else "TBA", font_data, "left"),
            (s.room.name, font_data, "left"),
            (s.get_session_type_display(), font_data, "center"),
            (s.get_status_display(), font_data, "center"),
        ]
        for col_idx, (val, font, align_h) in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = font
            cell.fill = row_fill
            cell.alignment = Alignment(horizontal=align_h, vertical="center")
            cell.border = thin_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
    ws.column_dimensions["F"].width = 28
    ws.freeze_panes = "A5"

    output = io.BytesIO()
    finish_worksheet(ws, header_row=4)
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_timetable_pdf(queryset, site_name="University Management System", logo_path=None, filter_text=None):
    buffer = io.BytesIO()
    doc = ReportDocTemplate(buffer, pagesize=landscape(A4),
                            leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = document_styles()

    title_style = ParagraphStyle("TTTitle", parent=styles["Normal"], fontName="Quicksand-Bold",
                                 fontSize=15, leading=19, textColor=colors.HexColor("#1E293B"))
    meta_style = ParagraphStyle("TTMeta", parent=styles["Normal"], fontName="Quicksand",
                                fontSize=9, leading=13, textColor=colors.HexColor("#64748B"), alignment=2)
    th_style = ParagraphStyle("TTTH", parent=styles["Normal"], fontName="Quicksand-Bold",
                              fontSize=8, leading=10, textColor=colors.white)
    td_style = ParagraphStyle("TTTD", parent=styles["Normal"], fontName="Quicksand",
                              fontSize=8, leading=10, textColor=colors.HexColor("#1E293B"))
    td_bold = ParagraphStyle("TTTDBold", parent=styles["Normal"], fontName="Quicksand-Bold",
                             fontSize=8, leading=10, textColor=colors.HexColor("#6C5CE7"))
    td_center = ParagraphStyle("TTTDCenter", parent=td_style, alignment=1)

    story = []
    branding = get_branding()
    branding.update(site_name=site_name, logo_path=logo_path or branding['logo_path'])
    story.append(letterhead(doc.width, 'Timetable Report',
        subtitle=f"Total records: {queryset.count()}", filter_text=filter_text, branding=branding))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#6C5CE7"), spaceAfter=12))

    headers = [Paragraph(f"<b>{h}</b>", th_style) for h in
              ["Day", "Time", "Course", "Lecturer", "Venue", "Programme"]]
    col_widths = [65, 90, 220, 150, 100, 100]
    data = [headers]

    for s in queryset:
        faculty_name = s.course.faculty.user.display_name if s.course.faculty else "TBA"
        prog_code = s.course.program.code if s.course.program else "—"
        data.append([
            Paragraph(s.get_day_display(), td_bold),
            Paragraph(f"{s.start_time.strftime('%H:%M')}–{s.end_time.strftime('%H:%M')}", td_style),
            Paragraph(f"{s.course.code} — {s.course.title}", td_style),
            Paragraph(faculty_name, td_style),
            Paragraph(s.room.name, td_center),
            Paragraph(prog_code, td_center),
        ])

    if len(data) == 1:
        data.append([Paragraph("<i>No timetable entries match the active filters.</i>", td_style)] + [""] * 5)

    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6C5CE7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    story.append(t)
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# ==============================================================================
# 2. TEMPLATE GENERATION (CSV & EXCEL)
# ==============================================================================

def generate_timetable_template_csv():
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    headers = [col[1] for col in TIMETABLE_IMPORT_COLUMNS]
    writer.writerow(headers)
    writer.writerow([SAMPLE_TIMETABLE_ROW[key] for key, _, _ in TIMETABLE_IMPORT_COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


def generate_timetable_template_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable Import Template"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    sample_font = Font(name="Arial", size=10, color="334155")
    sample_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border = Border(left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"),
                    top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1"))

    headers = [col[1] for col in TIMETABLE_IMPORT_COLUMNS]
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    sample_values = [SAMPLE_TIMETABLE_ROW[key] for key, _, _ in TIMETABLE_IMPORT_COLUMNS]
    ws.append(sample_values)
    ws.row_dimensions[2].height = 22
    for col_idx in range(1, len(sample_values) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = sample_font
        cell.fill = sample_fill
        cell.alignment = Alignment(vertical="center")
        cell.border = border

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==============================================================================
# 3. PARSING, VALIDATION & BULK IMPORT ENGINE
# ==============================================================================

def _normalize_timetable_header(header_name):
    if not header_name:
        return ""
    s = re.sub(r"[^a-zA-Z0-9]+", "_", str(header_name).strip().lower()).strip("_")
    alias_map = {
        "academic_year_term": "term_name", "academic_year": "term_name", "term": "term_name",
        "year_term": "term_name", "semester": "term_name",
        "day": "day", "day_of_week": "day",
        "start_time": "start_time", "start": "start_time",
        "end_time": "end_time", "end": "end_time",
        "course_code": "course_code", "code": "course_code", "subject_code": "course_code",
        "course_name": "course_name", "course_title": "course_name", "title": "course_name",
        "department": "department_code", "department_code": "department_code", "dept": "department_code",
        "programme": "programme_code", "program": "programme_code",
        "programme_code": "programme_code", "program_code": "programme_code",
        "class_year": "class_year", "class": "class_year", "year": "class_year",
        "faculty": "faculty", "lecturer": "faculty", "faculty_lecturer": "faculty",
        "instructor": "faculty", "teacher": "faculty",
        "venue": "room_name", "venue_room": "room_name", "room": "room_name", "room_name": "room_name",
        "session_type": "session_type", "type": "session_type", "class_type": "session_type",
    }
    return alias_map.get(s, s)


def parse_uploaded_timetable_file(uploaded_file):
    filename = (getattr(uploaded_file, "name", "") or "upload.csv").lower()
    rows = []

    if filename.endswith(".csv"):
        raw_bytes = uploaded_file.read()
        try:
            text = raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1")
        reader = csv.reader(io.StringIO(text))
        header_row = None
        for r in reader:
            if not any(str(c).strip() for c in r):
                continue
            if header_row is None:
                header_row = [_normalize_timetable_header(c) for c in r]
                continue
            row_dict = {header_row[i]: (r[i].strip() if i < len(r) else "") for i in range(len(header_row))}
            rows.append(row_dict)

    elif filename.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        header_row = None
        for r in ws.iter_rows(values_only=True):
            if not any(r):
                continue
            if header_row is None:
                header_row = [_normalize_timetable_header(c) for c in r]
                continue
            row_dict = {}
            for idx, col_key in enumerate(header_row):
                val = r[idx] if idx < len(r) else None
                if val is not None and hasattr(val, "strftime"):
                    val = val.strftime("%H:%M") if hasattr(val, "hour") and not hasattr(val, "year") else str(val)
                row_dict[col_key] = str(val).strip() if val is not None else ""
            rows.append(row_dict)

    return rows


def _parse_day(raw):
    if not raw:
        return None
    s = raw.strip().upper()[:3]
    return s if s in DAY_ORDER else None


def _parse_time(raw):
    if not raw:
        return None
    raw = raw.strip()
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p"):
        try:
            return datetime.strptime(raw, fmt).time()
        except ValueError:
            continue
    return None


def validate_timetable_import_rows(raw_rows):
    """
    Validate parsed timetable rows before importing.
    Row status: 'error' (missing/invalid data) > 'duplicate' (identical entry
    already exists) > 'conflict' (room, faculty or cohort double-booking) > 'valid'.
    """
    terms_by_name = {t.name.lower(): t for t in AcademicTerm.objects.all()}
    courses_by_code = {c.code.upper(): c for c in
                       Course.objects.select_related("department", "program", "faculty__user").all()}
    rooms_by_name = {r.name.lower(): r for r in ExamRoom.objects.all()}
    faculty_by_eid = {f.employee_id.upper(): f for f in FacultyProfile.objects.select_related("user").all()}

    session_type_map = {v.lower(): k for k, v in ClassSchedule.SessionType.choices}
    session_type_map.update({k.lower(): k for k in ClassSchedule.SessionType.values})

    # Existing DB entries, used for duplicate/conflict lookups.
    existing = [{
        "term_id": s.term_id, "day": s.day, "start": s.start_time, "end": s.end_time,
        "room_id": s.room_id, "course_id": s.course_id,
        "faculty_id": s.course.faculty_id, "program_id": s.course.program_id,
        "semester_no": s.course.semester_no, "code": s.course.code,
    } for s in ClassSchedule.objects.select_related("course")]

    accepted = list(existing)
    validated_items, valid_count, duplicate_count, conflict_count, error_count = [], 0, 0, 0, 0

    for idx, raw in enumerate(raw_rows, start=2):
        errors = []

        term_raw = raw.get("term_name", "").strip()
        day_raw = raw.get("day", "").strip()
        start_raw = raw.get("start_time", "").strip()
        end_raw = raw.get("end_time", "").strip()
        course_raw = raw.get("course_code", "").strip().upper()
        room_raw = raw.get("room_name", "").strip()
        session_raw = raw.get("session_type", "").strip()
        faculty_raw = raw.get("faculty", "").strip()

        if not term_raw:
            errors.append("Academic Year / Term is required.")
        term_obj = terms_by_name.get(term_raw.lower()) if term_raw else None
        if term_raw and not term_obj:
            errors.append(f"Unknown academic term '{term_raw}'.")

        if not day_raw:
            errors.append("Day is required.")
        day_code = _parse_day(day_raw)
        if day_raw and not day_code:
            errors.append(f"Invalid day '{day_raw}'. Use Monday–Saturday.")

        if not start_raw:
            errors.append("Start Time is required.")
        start_t = _parse_time(start_raw)
        if start_raw and not start_t:
            errors.append(f"Invalid start time '{start_raw}'. Use HH:MM (24h).")

        if not end_raw:
            errors.append("End Time is required.")
        end_t = _parse_time(end_raw)
        if end_raw and not end_t:
            errors.append(f"Invalid end time '{end_raw}'. Use HH:MM (24h).")
        if start_t and end_t and end_t <= start_t:
            errors.append("End Time must be after Start Time.")

        if not course_raw:
            errors.append("Course Code is required.")
        course_obj = courses_by_code.get(course_raw) if course_raw else None
        if course_raw and not course_obj:
            errors.append(f"Unknown course code '{course_raw}'.")

        if not room_raw:
            errors.append("Venue / Room is required.")
        room_obj = rooms_by_name.get(room_raw.lower()) if room_raw else None
        if room_raw and not room_obj:
            errors.append(f"Unknown venue / room '{room_raw}'.")

        session_type = session_type_map.get(session_raw.lower(), ClassSchedule.SessionType.LECTURE)

        faculty_obj = None
        if faculty_raw:
            faculty_obj = faculty_by_eid.get(faculty_raw.upper())
            if not faculty_obj:
                matched = [f for f in faculty_by_eid.values() if faculty_raw.lower() in f.user.display_name.lower()]
                faculty_obj = matched[0] if matched else None

        if errors:
            error_count += 1
            validated_items.append({
                "row_num": idx, "status": "error", "term_name": term_raw, "day": day_raw,
                "start_time": start_raw, "end_time": end_raw, "course_code": course_raw,
                "course_name": course_obj.title if course_obj else "", "room_name": room_raw,
                "session_type": ClassSchedule.SessionType.LECTURE, "faculty_name": faculty_raw or "TBA",
                "term_id": None, "course_id": None, "room_id": None, "errors": errors,
            })
            continue

        effective_faculty_id = course_obj.faculty_id or (faculty_obj.id if faculty_obj else None)

        is_duplicate = any(
            e["term_id"] == term_obj.id and e["course_id"] == course_obj.id and e["day"] == day_code
            and e["start"] == start_t and e["end"] == end_t and e["room_id"] == room_obj.id
            for e in accepted)

        conflict_reasons = []
        if not is_duplicate:
            overlapping = [e for e in accepted if e["term_id"] == term_obj.id and e["day"] == day_code
                          and e["start"] < end_t and e["end"] > start_t and e["course_id"] != course_obj.id]
            room_clash = next((e for e in overlapping if e["room_id"] == room_obj.id), None)
            if room_clash:
                conflict_reasons.append(f"Venue '{room_obj.name}' already booked for {room_clash['code']} at that time.")
            if effective_faculty_id:
                faculty_clash = next((e for e in overlapping if e["faculty_id"] == effective_faculty_id), None)
                if faculty_clash:
                    conflict_reasons.append(f"Lecturer already teaching {faculty_clash['code']} at that time.")
            if course_obj.program_id:
                cohort_clash = next((e for e in overlapping if e["program_id"] == course_obj.program_id
                                    and e["semester_no"] == course_obj.semester_no), None)
                if cohort_clash:
                    conflict_reasons.append(
                        f"{course_obj.program} Sem {course_obj.semester_no} already has "
                        f"{cohort_clash['code']} scheduled at that time.")

        if is_duplicate:
            status = "duplicate"
            duplicate_count += 1
            errors.append("An identical timetable entry already exists.")
        elif conflict_reasons:
            status = "conflict"
            conflict_count += 1
            errors.extend(conflict_reasons)
        else:
            status = "valid"
            valid_count += 1
            accepted.append({
                "term_id": term_obj.id, "day": day_code, "start": start_t, "end": end_t,
                "room_id": room_obj.id, "course_id": course_obj.id,
                "faculty_id": effective_faculty_id, "program_id": course_obj.program_id,
                "semester_no": course_obj.semester_no, "code": course_obj.code,
            })

        validated_items.append({
            "row_num": idx, "status": status, "term_name": term_obj.name, "day": day_code,
            "day_display": DAY_LABELS.get(day_code, day_code),
            "start_time": start_t.strftime("%H:%M"), "end_time": end_t.strftime("%H:%M"),
            "course_code": course_obj.code, "course_name": course_obj.title, "room_name": room_obj.name,
            "session_type": session_type,
            "faculty_name": course_obj.faculty.user.display_name if course_obj.faculty else (
                faculty_obj.user.display_name if faculty_obj else "TBA"),
            "term_id": term_obj.id, "course_id": course_obj.id, "room_id": room_obj.id, "errors": errors,
        })

    return {
        "items": validated_items, "total_count": len(raw_rows), "valid_count": valid_count,
        "duplicate_count": duplicate_count, "conflict_count": conflict_count, "error_count": error_count,
    }


def execute_timetable_import(valid_items, status=ClassSchedule.Status.DRAFT):
    imported_count, failed_count = 0, 0
    with transaction.atomic():
        for item in valid_items:
            try:
                ClassSchedule.objects.create(
                    term_id=item["term_id"], course_id=item["course_id"], room_id=item["room_id"],
                    day=item["day"], start_time=_parse_time(item["start_time"]),
                    end_time=_parse_time(item["end_time"]), session_type=item["session_type"],
                    status=status)
                imported_count += 1
            except Exception:
                failed_count += 1
    return imported_count, failed_count
