from university.document_design import ReportDocTemplate, document_styles, PageNumberCanvas, letterhead, get_branding, finish_worksheet
import csv
import io
import os
from decimal import Decimal

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.signals import (
    user_logged_in, user_logged_out, user_login_failed
)
from django.dispatch import receiver
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable, Image as RLImage, Paragraph, SimpleDocTemplate,
    Spacer, Table, TableStyle
)

from university.models import AuditLog

User = get_user_model()


def get_client_ip(request):
    """Extract client IP honoring proxy forward headers."""
    if not request:
        return None
    forwarded = request.META.get("HTTP_X_FORWARDED_FOR", "")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.META.get("REMOTE_ADDR") or None


def detect_device_type(user_agent_str):
    """Categorize User-Agent header into Desktop, Mobile, Tablet, or Bot."""
    if not user_agent_str:
        return "Desktop"
    ua = user_agent_str.lower()
    if any(bot in ua for bot in ["bot", "spider", "crawler", "curl", "python-requests"]):
        return "Bot / API"
    if "tablet" in ua or "ipad" in ua:
        return "Tablet"
    if any(m in ua for m in ["mobi", "iphone", "android", "blackberry", "windows phone"]):
        return "Mobile"
    return "Desktop"


def log_activity(
    request=None,
    user=None,
    action=AuditLog.Action.UPDATE,
    module=AuditLog.Module.ACADEMICS,
    entity="",
    entity_id="",
    description="",
    previous_state=None,
    new_state=None,
):
    """
    Central tamper-resistant audit logging function.
    Captures actor, action, module, diff JSON, IP address, and device type.
    """
    actor = user
    if not actor and request and hasattr(request, "user") and request.user.is_authenticated:
        actor = request.user

    user_disp = "System"
    user_role = "System"
    if actor and actor.is_authenticated:
        user_disp = getattr(actor, "display_name", str(actor))
        user_role = getattr(actor, "get_role_display", lambda: getattr(actor, "role", "User"))()
        if getattr(actor, "is_superuser", False):
            user_role = "Superuser"
    elif actor:
        user_disp = str(actor)

    ip = get_client_ip(request) if request else None
    ua_str = request.META.get("HTTP_USER_AGENT", "") if request else ""
    device = detect_device_type(ua_str) if ua_str else "Desktop"

    return AuditLog.objects.create(
        user=actor if actor and actor.is_authenticated else None,
        user_display=user_disp,
        user_role=user_role,
        action=action,
        module=module,
        entity=entity,
        entity_id=str(entity_id) if entity_id is not None else "",
        description=description,
        previous_state=previous_state,
        new_state=new_state,
        ip_address=ip,
        user_agent=ua_str[:500] if ua_str else "",
        device_type=device,
    )


# ==============================================================================
# AUTHENTICATION SIGNALS
# ==============================================================================

@receiver(user_logged_in)
def audit_login_success(sender, request, user, **kwargs):
    log_activity(
        request=request,
        user=user,
        action=AuditLog.Action.LOGIN,
        module=AuditLog.Module.AUTH,
        entity="User",
        entity_id=user.pk,
        description=f"Successful login for user '{user.username}' ({user.get_role_display() if hasattr(user, 'get_role_display') else 'User'}).",
    )


@receiver(user_logged_out)
def audit_logout(sender, request, user, **kwargs):
    if user and user.is_authenticated:
        log_activity(
            request=request,
            user=user,
            action=AuditLog.Action.LOGOUT,
            module=AuditLog.Module.AUTH,
            entity="User",
            entity_id=user.pk,
            description=f"User '{user.username}' logged out.",
        )


@receiver(user_login_failed)
def audit_login_failed(sender, credentials, request, **kwargs):
    username = credentials.get("username", "Unknown")
    log_activity(
        request=request,
        user=None,
        action=AuditLog.Action.FAILED_LOGIN,
        module=AuditLog.Module.AUTH,
        entity="User",
        entity_id=username,
        description=f"Failed login attempt for username '{username}'.",
    )


# ==============================================================================
# AUDIT LOG EXPORT GENERATORS (CSV, EXCEL, PDF)
# ==============================================================================

def export_audit_csv(queryset):
    """Export audit log queryset as CSV."""
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow([
        "Timestamp", "User", "Role", "Action", "Module",
        "Entity", "Entity ID", "Description", "IP Address", "Device"
    ])
    for log in queryset:
        writer.writerow([
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            log.user_display,
            log.user_role,
            log.get_action_display() if hasattr(log, "get_action_display") else log.action,
            log.get_module_display() if hasattr(log, "get_module_display") else log.module,
            log.entity,
            log.entity_id,
            log.description,
            log.ip_address or "—",
            log.device_type,
        ])
    return output.getvalue().encode("utf-8")


def export_audit_excel(queryset, site_name="University Management System"):
    """Export audit log queryset as Excel (.xlsx)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Trails"

    # Header styling
    brand_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    brand_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )

    headers = [
        "Timestamp", "User", "Role", "Action", "Module",
        "Entity", "Entity ID", "Description", "IP Address", "Device"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = brand_fill
        cell.font = brand_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for log in queryset:
        ws.append([
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            log.user_display,
            log.user_role,
            log.action,
            log.module,
            log.entity,
            log.entity_id,
            log.description,
            log.ip_address or "—",
            log.device_type,
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = io.BytesIO()
    finish_worksheet(ws, header_row=1)
    wb.save(stream)
    return stream.getvalue()


def export_audit_pdf(queryset, site_name="University Management System"):
    """Export audit log queryset as landscape PDF document."""
    buffer = io.BytesIO()
    doc = ReportDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=25,
        rightMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = document_styles()
    primary_color = colors.HexColor("#1e3a8a")
    dark_gray = colors.HexColor("#1f2937")

    title_style = ParagraphStyle(
        "AuditTitle",
        parent=styles["Normal"],
        fontName="Quicksand-Bold",
        fontSize=15,
        leading=19,
        textColor=primary_color,
        alignment=1,
    )
    body_style = ParagraphStyle(
        "AuditBody",
        parent=styles["Normal"],
        fontName="Quicksand",
        fontSize=8,
        leading=11,
        textColor=dark_gray,
    )
    body_bold = ParagraphStyle(
        "AuditBodyBold",
        parent=styles["Normal"],
        fontName="Quicksand-Bold",
        fontSize=8,
        leading=11,
        textColor=dark_gray,
    )

    story = []
    story.append(letterhead(doc.width, "System Audit Trails & Activity Log"))
    story.append(Paragraph(f"Generated on {timezone.now().strftime('%d %B %Y at %H:%M:%S')} · Official Immutable System Log", ParagraphStyle("Sub", parent=body_style, alignment=1, textColor=colors.HexColor("#6b7280"))))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceAfter=10))

    table_data = [
        [
            Paragraph("<b>Timestamp</b>", body_bold),
            Paragraph("<b>User & Role</b>", body_bold),
            Paragraph("<b>Action</b>", body_bold),
            Paragraph("<b>Module & Entity</b>", body_bold),
            Paragraph("<b>Description</b>", body_bold),
            Paragraph("<b>IP & Device</b>", body_bold),
        ]
    ]

    for log in queryset[:200]:  # Limit to 200 rows per PDF export for fast rendering
        table_data.append([
            Paragraph(log.timestamp.strftime("%Y-%m-%d<br/>%H:%M:%S"), body_style),
            Paragraph(f"<b>{log.user_display}</b><br/>{log.user_role}", body_style),
            Paragraph(f"<b>{log.action}</b>", body_style),
            Paragraph(f"<b>{log.module}</b><br/>{log.entity} #{log.entity_id}", body_style),
            Paragraph(log.description[:120], body_style),
            Paragraph(f"{log.ip_address or '—'}<br/>{log.device_type}", body_style),
        ])

    col_widths = [doc.width * fraction for fraction in [.10, .14, .12, .16, .32, .16]]
    audit_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    audit_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(audit_table)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
