import csv
import io
import re
from datetime import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from university.document_design import (ReportDocTemplate, document_styles, document_fonts,
    PageNumberCanvas, letterhead, get_branding, finish_worksheet)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable, Image as RLImage, KeepTogether, Paragraph, SimpleDocTemplate,
    Spacer, Table, TableStyle
)

from accounts.models import Role, StudentProfile
from university.models import Program

User = get_user_model()

# Standard headers for import & export
IMPORT_COLUMNS = [
    ("roll_no", "Roll No", True),
    ("first_name", "First Name", True),
    ("last_name", "Last Name", True),
    ("email", "Email", True),
    ("phone", "Phone", False),
    ("username", "Username", False),
    ("password", "Password", False),
    ("program_code", "Program Code", False),
    ("current_semester", "Current Semester", False),
    ("gender", "Gender", False),
    ("address", "Address", False),
    ("guardian_name", "Guardian Name", False),
]

SAMPLE_STUDENT_ROW = {
    "roll_no": "STU2026101",
    "first_name": "Jane",
    "last_name": "Doe",
    "email": "jane.doe@example.com",
    "phone": "+254 712 345 678",
    "username": "jane.doe",
    "password": "Student@2026",
    "program_code": "BT-CSE",
    "current_semester": "1",
    "gender": "Female",
    "address": "Hostel Block B, Room 204",
    "guardian_name": "Robert Doe",
}


# ==============================================================================
# 1. EXPORT FUNCTIONS
# ==============================================================================

def export_students_csv(queryset):
    """Generate CSV byte string with UTF-8 BOM."""
    buffer = io.StringIO()
    # Write UTF-8 BOM so Excel opens it with correct encoding
    buffer.write("\ufeff")
    writer = csv.writer(buffer)

    headers = [
        "Roll No", "First Name", "Last Name", "Full Name", "Email", "Phone",
        "Program Code", "Program Name", "Semester", "Gender", "Address",
        "Guardian Name", "Admission Date"
    ]
    writer.writerow(headers)

    for sp in queryset:
        u = sp.user
        prog_code = sp.program.code if sp.program else "—"
        prog_name = sp.program.name if sp.program else "—"
        adm_date = sp.admission_date.strftime("%Y-%m-%d") if sp.admission_date else "—"
        writer.writerow([
            sp.roll_no,
            u.first_name,
            u.last_name,
            u.display_name,
            u.email,
            u.phone,
            prog_code,
            prog_name,
            f"Sem {sp.current_semester}",
            sp.get_gender_display(),
            sp.address,
            sp.guardian_name,
            adm_date,
        ])

    return buffer.getvalue().encode("utf-8")


def export_students_excel(queryset, site_name="University Management System"):
    """Generate professional styled Excel (.xlsx) file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.views.sheetView[0].showGridLines = True

    # Color definitions (UMS palette)
    primary_color = "6C5CE7"      # Brand Purple
    header_fill_color = "4834D4"  # Darker purple for header
    zebra_color = "F7F8FC"        # Soft surface tint
    border_color = "D6D9E6"

    font_title = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=10, italic=True, color="E8E6FF")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10, color="2B2B3A")
    font_bold_data = Font(name="Arial", size=10, bold=True, color="2B2B3A")

    thin_border = Border(
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )

    # 1. Title Banner
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"  {site_name.upper()} — STUDENT RECORDS"
    title_cell.font = font_title
    title_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 36

    # 2. Subtitle / Metadata Row
    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    sub_cell.value = f"  Generated: {timezone.now().strftime('%d %b %Y, %H:%M')} | Total Students: {queryset.count()}"
    sub_cell.font = font_sub
    sub_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    sub_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 20

    # Row 3 is spacer
    ws.row_dimensions[3].height = 8

    # 3. Column Headers
    headers = [
        "Roll No", "First Name", "Last Name", "Email", "Phone",
        "Program", "Semester", "Gender", "Address", "Guardian Name", "Admission Date"
    ]
    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 26

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 7, 8, 11] else "left", vertical="center")
        cell.border = thin_border

    # 4. Data Rows
    for row_idx, sp in enumerate(queryset, start=5):
        u = sp.user
        prog_str = f"{sp.program.code} - {sp.program.name}" if sp.program else "—"
        adm_date = sp.admission_date.strftime("%Y-%m-%d") if sp.admission_date else "—"
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color=zebra_color, end_color=zebra_color, fill_type="solid") if is_even else None

        values = [
            sp.roll_no,
            u.first_name,
            u.last_name,
            u.email,
            u.phone,
            prog_str,
            f"Sem {sp.current_semester}",
            sp.get_gender_display(),
            sp.address,
            sp.guardian_name,
            adm_date,
        ]

        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_bold_data if col_idx in [1, 2] else font_data
            if row_fill:
                cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center" if col_idx in [1, 7, 8, 11] else "left",
                vertical="center"
            )

    # 5. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip merged title rows for column width measurement
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze header rows
    ws.freeze_panes = "A5"

    output = io.BytesIO()
    finish_worksheet(ws, header_row=4)
    wb.save(output)
    output.seek(0)
    return output.getvalue()


NumberedCanvas = PageNumberCanvas


def export_students_pdf(queryset, site_name="University Management System", logo_path=None, filter_text=None):
    """Generate high-quality branded PDF report in landscape A4."""
    buffer = io.BytesIO()
    doc = ReportDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = document_styles()
    
    # Custom Typography Styles
    title_style = ParagraphStyle(
        "ReportTitle",
        fontName="Quicksand-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#2B2B3A"),
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        fontName="Quicksand",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#718096"),
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        fontName="Quicksand-Bold",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#6C5CE7"),
        alignment=2,  # Right aligned
    )
    th_style = ParagraphStyle(
        "TableHeader",
        fontName="Quicksand-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=0,
    )
    td_style = ParagraphStyle(
        "TableCell",
        fontName="Quicksand",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#2B2B3A"),
    )
    td_bold = ParagraphStyle(
        "TableCellBold",
        fontName="Quicksand-Bold",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#2B2B3A"),
    )
    td_center = ParagraphStyle(
        "TableCellCenter",
        fontName="Quicksand",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#2B2B3A"),
        alignment=1,
    )

    story = []

    branding = get_branding()
    branding.update(site_name=site_name, logo_path=logo_path or branding['logo_path'])
    story.append(letterhead(doc.width, 'Student Directory',
        subtitle=f"Total records: {queryset.count()}", filter_text=filter_text, branding=branding))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#6C5CE7"), spaceAfter=12))

    # 2. Table of Students
    table_headers = [
        Paragraph("Roll No", th_style),
        Paragraph("Full Name", th_style),
        Paragraph("Program", th_style),
        Paragraph("Sem", th_style),
        Paragraph("Gender", th_style),
        Paragraph("Email Address", th_style),
        Paragraph("Phone", th_style),
        Paragraph("Guardian", th_style),
    ]
    table_data = [table_headers]

    for sp in queryset:
        u = sp.user
        prog_code = sp.program.code if sp.program else "—"
        table_data.append([
            Paragraph(sp.roll_no, td_bold),
            Paragraph(u.display_name, td_style),
            Paragraph(prog_code, td_center),
            Paragraph(f"{sp.current_semester}", td_center),
            Paragraph(sp.get_gender_display(), td_center),
            Paragraph(u.email, td_style),
            Paragraph(u.phone or "—", td_style),
            Paragraph(sp.guardian_name or "—", td_style),
        ])

    # Table Widths: Total printable landscape A4 is ~770pt
    # 70 + 130 + 65 + 35 + 50 + 190 + 95 + 135 = 770
    col_widths = [70, 130, 65, 35, 50, 190, 95, 135]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6C5CE7")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
    ]

    # Alternating row background
    for r in range(1, len(table_data)):
        if r % 2 == 0:
            t_style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F8F9FE")))

    t.setStyle(TableStyle(t_style))
    story.append(t)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# ==============================================================================
# 2. TEMPLATE GENERATION (CSV & EXCEL)
# ==============================================================================

def generate_import_template_csv():
    """Generate CSV template with Row 1 headers and Row 2 realistic sample student data."""
    buffer = io.StringIO()
    buffer.write("\ufeff")
    writer = csv.writer(buffer)
    
    headers = [col[1] for col in IMPORT_COLUMNS]
    writer.writerow(headers)
    
    sample_row = [
        SAMPLE_STUDENT_ROW["roll_no"],
        SAMPLE_STUDENT_ROW["first_name"],
        SAMPLE_STUDENT_ROW["last_name"],
        SAMPLE_STUDENT_ROW["email"],
        SAMPLE_STUDENT_ROW["phone"],
        SAMPLE_STUDENT_ROW["username"],
        SAMPLE_STUDENT_ROW["password"],
        SAMPLE_STUDENT_ROW["program_code"],
        SAMPLE_STUDENT_ROW["current_semester"],
        SAMPLE_STUDENT_ROW["gender"],
        SAMPLE_STUDENT_ROW["address"],
        SAMPLE_STUDENT_ROW["guardian_name"],
    ]
    writer.writerow(sample_row)
    return buffer.getvalue().encode("utf-8")


def generate_import_template_excel():
    """Generate Excel (.xlsx) template with Row 1 headers and Row 2 prefilled sample data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Import Template"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
    sample_fill = PatternFill(start_color="F4F5FB", end_color="F4F5FB", fill_type="solid")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_sample = Font(name="Arial", size=10, italic=True, color="4834D4")
    
    thin_border = Border(
        left=Side(style="thin", color="D6D9E6"),
        right=Side(style="thin", color="D6D9E6"),
        top=Side(style="thin", color="D6D9E6"),
        bottom=Side(style="thin", color="D6D9E6"),
    )

    headers = [col[1] for col in IMPORT_COLUMNS]
    ws.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sample_values = [
        SAMPLE_STUDENT_ROW["roll_no"],
        SAMPLE_STUDENT_ROW["first_name"],
        SAMPLE_STUDENT_ROW["last_name"],
        SAMPLE_STUDENT_ROW["email"],
        SAMPLE_STUDENT_ROW["phone"],
        SAMPLE_STUDENT_ROW["username"],
        SAMPLE_STUDENT_ROW["password"],
        SAMPLE_STUDENT_ROW["program_code"],
        int(SAMPLE_STUDENT_ROW["current_semester"]),
        SAMPLE_STUDENT_ROW["gender"],
        SAMPLE_STUDENT_ROW["address"],
        SAMPLE_STUDENT_ROW["guardian_name"],
    ]
    ws.row_dimensions[2].height = 22
    for col_idx, val in enumerate(sample_values, start=1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.font = font_sample
        cell.fill = sample_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==============================================================================
# 3. IMPORT PARSER & VALIDATOR
# ==============================================================================

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

def _normalize_header(name):
    """Normalize column header: 'Roll No.' -> 'roll_no'."""
    s = re.sub(r"[^a-zA-Z0-9]+", "_", str(name).strip().lower()).strip("_")
    alias_map = {
        "roll": "roll_no",
        "roll_number": "roll_no",
        "reg_no": "roll_no",
        "registration_no": "roll_no",
        "student_id": "roll_no",
        "fname": "first_name",
        "lname": "last_name",
        "mail": "email",
        "e_mail": "email",
        "tel": "phone",
        "phone_number": "phone",
        "mobile": "phone",
        "user_name": "username",
        "program": "program_code",
        "course": "program_code",
        "course_code": "program_code",
        "sem": "current_semester",
        "semester": "current_semester",
        "sex": "gender",
        "guardian": "guardian_name",
        "parent_name": "guardian_name",
    }
    return alias_map.get(s, s)


def parse_uploaded_file(uploaded_file):
    """Read CSV or Excel rows as a list of dicts."""
    filename = uploaded_file.name.lower()
    rows = []

    if filename.endswith(".csv"):
        raw_bytes = uploaded_file.read()
        # Decode UTF-8 or Latin-1
        try:
            text = raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1")
        
        reader = csv.reader(io.StringIO(text))
        header_row = None
        for r_idx, row in enumerate(reader):
            if not any(str(c).strip() for c in row):
                continue
            if header_row is None:
                header_row = [_normalize_header(c) for c in row]
                continue
            row_dict = {}
            for h, val in zip(header_row, row):
                if h:
                    row_dict[h] = str(val).strip()
            rows.append(row_dict)

    elif filename.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        header_row = None
        for r in ws.iter_rows(values_only=True):
            if not any(c is not None and str(c).strip() for c in r):
                continue
            if header_row is None:
                header_row = [_normalize_header(c) for c in r]
                continue
            row_dict = {}
            for h, val in zip(header_row, r):
                if h:
                    row_dict[h] = "" if val is None else str(val).strip()
            rows.append(row_dict)
    else:
        raise ValueError("Unsupported file format. Please upload a .csv or .xlsx file.")

    return rows


def validate_import_rows(raw_rows):
    """
    Validate imported student rows, detect file and DB duplicates,
    and return structured result for user preview before commit.
    """
    existing_rolls = set(StudentProfile.objects.values_list("roll_no", flat=True))
    existing_usernames = set(User.objects.values_list("username", flat=True))
    existing_emails = set(User.objects.values_list("email", flat=True))

    programs_by_code = {p.code.upper(): p for p in Program.objects.all() if p.code}
    programs_by_name = {p.name.strip().lower(): p for p in Program.objects.all() if p.name}

    seen_rolls = set()
    seen_usernames = set()
    seen_emails = set()

    validated_items = []
    valid_count = 0
    duplicate_count = 0
    error_count = 0

    for idx, raw in enumerate(raw_rows, start=2):  # Row 1 is header
        errors = []
        is_duplicate = False

        # 1. Required Fields Check
        roll_no = raw.get("roll_no", "").strip()
        first_name = raw.get("first_name", "").strip()
        last_name = raw.get("last_name", "").strip()
        email = raw.get("email", "").strip()
        phone = raw.get("phone", "0000").strip() or "0000"
        
        username = raw.get("username", "").strip()
        if not username and email:
            username = email.split("@")[0]
        if not username:
            username = f"{first_name.lower()}.{last_name.lower()}" if first_name else f"student_{idx}"
        username = re.sub(r"[^a-zA-Z0-9._-]", "", username)

        password = raw.get("password", "").strip() or "Student@2026"
        program_code = raw.get("program_code", "").strip().upper()
        
        raw_sem = raw.get("current_semester", "1").strip()
        gender_raw = raw.get("gender", "Other").strip().upper()
        address = raw.get("address", "").strip() or "Campus Hostel Block"
        guardian_name = raw.get("guardian_name", "").strip() or "Guardian"

        if not roll_no:
            errors.append("Roll number is required.")
        if not first_name:
            errors.append("First name is required.")
        if not last_name:
            errors.append("Last name is required.")
        if not email:
            errors.append("Email address is required.")
        elif not EMAIL_REGEX.match(email):
            errors.append(f"Invalid email format '{email}'.")

        # 2. Semester validation
        try:
            semester = int(raw_sem)
            if not (1 <= semester <= 12):
                errors.append(f"Semester must be between 1 and 12 (got '{raw_sem}').")
        except ValueError:
            errors.append(f"Invalid semester number '{raw_sem}'.")
            semester = 1

        # 3. Gender validation
        gender_choice = "O"
        if gender_raw in ["M", "MALE"]:
            gender_choice = "M"
        elif gender_raw in ["F", "FEMALE"]:
            gender_choice = "F"

        program_obj = None
        if program_code:
            program_obj = programs_by_code.get(program_code) or programs_by_name.get(program_code.lower())
            if not program_obj:
                errors.append(f"Unknown program code '{program_code}'.")

        # 5. Duplicate Detection (In-file and Database)
        dup_reasons = []
        if roll_no:
            if roll_no in seen_rolls:
                dup_reasons.append(f"Duplicate Roll No '{roll_no}' in this file")
            elif roll_no in existing_rolls:
                dup_reasons.append(f"Roll No '{roll_no}' already exists in database")

        if username:
            if username in seen_usernames:
                dup_reasons.append(f"Duplicate Username '{username}' in this file")
            elif username in existing_usernames:
                dup_reasons.append(f"Username '{username}' already taken in database")

        if email:
            if email in seen_emails:
                dup_reasons.append(f"Duplicate Email '{email}' in this file")
            elif email in existing_emails:
                dup_reasons.append(f"Email '{email}' already registered in database")

        if dup_reasons:
            is_duplicate = True
            errors.extend(dup_reasons)

        # Track seen identifiers for intra-file collision check
        if roll_no:
            seen_rolls.add(roll_no)
        if username:
            seen_usernames.add(username)
        if email:
            seen_emails.add(email)

        # Determine Row Status
        if errors and not is_duplicate:
            status = "error"
            error_count += 1
        elif is_duplicate:
            status = "duplicate"
            duplicate_count += 1
        else:
            status = "valid"
            valid_count += 1

        item = {
            "row_num": idx,
            "roll_no": roll_no,
            "first_name": first_name,
            "last_name": last_name,
            "full_name": f"{first_name} {last_name}".strip(),
            "email": email,
            "phone": phone,
            "username": username,
            "password": password,
            "program_code": program_code,
            "program_id": program_obj.id if program_obj else None,
            "program_name": program_obj.name if program_obj else "—",
            "current_semester": semester,
            "gender": gender_choice,
            "gender_display": "Male" if gender_choice == "M" else ("Female" if gender_choice == "F" else "Other"),
            "address": address,
            "guardian_name": guardian_name,
            "status": status,
            "errors": errors,
        }
        validated_items.append(item)

    return {
        "items": validated_items,
        "total_count": len(validated_items),
        "valid_count": valid_count,
        "duplicate_count": duplicate_count,
        "error_count": error_count,
    }


# ==============================================================================
# 4. BATCH EXECUTION
# ==============================================================================

def execute_student_import(valid_items):
    """
    Atomically insert validated student records and user accounts.
    Returns (imported_count, failed_count).
    """
    imported_count = 0
    failed_count = 0

    with transaction.atomic():
        for item in valid_items:
            try:
                # Double-check uniqueness before inserting
                if StudentProfile.objects.filter(roll_no=item["roll_no"]).exists():
                    failed_count += 1
                    continue
                if User.objects.filter(username=item["username"]).exists():
                    failed_count += 1
                    continue

                user = User(
                    username=item["username"],
                    email=item["email"],
                    first_name=item["first_name"],
                    last_name=item["last_name"],
                    phone=item["phone"],
                    role=Role.STUDENT,
                )
                user.set_password(item["password"])
                user.save()

                sp = StudentProfile(
                    user=user,
                    roll_no=item["roll_no"],
                    program_id=item["program_id"],
                    current_semester=item["current_semester"],
                    gender=item["gender"],
                    address=item["address"],
                    guardian_name=item["guardian_name"],
                    admission_date=timezone.now().date(),
                )
                sp.save()
                imported_count += 1
            except Exception:
                failed_count += 1

    return imported_count, failed_count
