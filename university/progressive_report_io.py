"""
Student Progressive Report Data Builder & Exporters (PDF, Excel, CSV)
Provides comprehensive 360-degree academic progress tracking and multi-format reports.
"""
from xml.sax.saxutils import escape
import io
import csv
from datetime import date
from decimal import Decimal
from collections import Counter

from django.utils import timezone
from django.conf import settings
import os

from university.document_design import (ReportDocTemplate, document_styles, document_fonts,
    PageNumberCanvas, letterhead, get_branding, finish_worksheet, make_qr_code_flowable)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import examination_services as workflow
from .transcript_io import build_transcript_context


def build_progressive_report_context(student):
    """
    Builds context dictionary for Student Progressive Report web view and file exports.
    """
    transcript_ctx = build_transcript_context(student)
    statement = workflow.student_statement(student)

    semesters_data = []
    gpa_trend_labels = []
    gpa_trend_values = []

    for idx, sem in enumerate(transcript_ctx.get("semesters", []), start=1):
        term_obj = sem.get("term")
        term_name = term_obj.name if term_obj else f"Semester {idx}"
        
        courses_list = []
        for g in sem.get("groups", []):
            components = g.get("components", [])
            ca_score = 0.0
            exam_score = 0.0
            for c in components:
                eff = c.get("effective")
                if eff and eff.marks_obtained is not None:
                    kind = eff.exam.kind
                    if kind in ["CAT", "ASSIGNMENT", "QUIZ", "MIDTERM"]:
                        ca_score += float(eff.marks_obtained)
                    else:
                        exam_score += float(eff.marks_obtained)

            total_pct = float(g.get("total") or 0.0)
            grade = g.get("grade") or "N/A"
            gp = float(g.get("grade_point") or 0.0)
            outcome = g.get("outcome") or "In Progress"

            courses_list.append({
                "code": g["course"].code,
                "title": g["course"].title,
                "credits": g["course"].credits,
                "ca_mark": round(ca_score, 1) if ca_score else "—",
                "exam_mark": round(exam_score, 1) if exam_score else "—",
                "total_pct": round(total_pct, 1),
                "grade": grade,
                "grade_point": gp,
                "outcome": outcome,
            })

        term_gpa = sem.get("term_gpa")
        cum_gpa = sem.get("cumulative_gpa")

        if term_gpa is not None:
            gpa_trend_labels.append(term_name)
            gpa_trend_values.append(float(term_gpa))

        semesters_data.append({
            "term_name": term_name,
            "term_obj": term_obj,
            "courses": courses_list,
            "credits_attempted": sem.get("credits_attempted", 0),
            "credits_completed": sem.get("credits_completed", 0),
            "term_gpa": round(float(term_gpa), 2) if term_gpa is not None else "N/A",
            "cumulative_gpa": round(float(cum_gpa), 2) if cum_gpa is not None else "N/A",
        })

    cgpa = transcript_ctx.get("cgpa")
    cgpa_formatted = round(float(cgpa), 2) if cgpa is not None else 0.0
    standing = transcript_ctx.get("overall_standing", "In Progress")

    req_credits = transcript_ctx.get("required_credits") or 120
    earned_credits = transcript_ctx.get("overall_credits_completed", 0)
    attempted_credits = transcript_ctx.get("overall_credits_attempted", 0)

    completion_pct = round((earned_credits / req_credits * 100), 1) if req_credits else 0.0
    pass_rate = round((earned_credits / max(attempted_credits, 1) * 100), 1) if attempted_credits else 0.0

    return {
        "student": student,
        "user": student.user,
        "program": student.program,
        "department": student.program.department if student.program else None,
        "semesters": semesters_data,
        "cgpa": cgpa_formatted,
        "overall_standing": standing,
        "required_credits": req_credits,
        "earned_credits": earned_credits,
        "attempted_credits": attempted_credits,
        "completion_pct": min(completion_pct, 100.0),
        "pass_rate": min(pass_rate, 100.0),
        "total_courses_completed": transcript_ctx.get("total_courses_completed", 0),
        "grade_distribution": transcript_ctx.get("grade_distribution", {}),
        "gpa_trend_labels": gpa_trend_labels,
        "gpa_trend_values": gpa_trend_values,
        "issued_date": timezone.localdate().strftime("%B %d, %Y"),
    }


def export_progressive_report_pdf(student, ctx, tracking_info=None, verify_url=None):
    """
    Generates formal A4 PDF Progressive Report via ReportLab with optional audit watermark & QR verification.
    """
    buf = io.BytesIO()
    doc = ReportDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = document_styles()
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Normal'],
        fontName='Quicksand-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E293B'),
        alignment=1
    )
    subtitle_style = ParagraphStyle(
        'RepSubTitle',
        parent=styles['Normal'],
        fontName='Quicksand-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#6C5CE7'),
        alignment=1
    )
    bold_style = ParagraphStyle(
        'RepBold',
        parent=styles['Normal'],
        fontName='Quicksand-Bold',
        fontSize=9,
        leading=11,
        keepWithNext=True,
        textColor=colors.HexColor('#1E293B')
    )
    normal_style = ParagraphStyle(
        'RepNormal',
        parent=styles['Normal'],
        fontName='Quicksand',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#333333')
    )
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Quicksand-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Quicksand',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )

    elements = []

    # Header
    elements.append(letterhead(doc.width, "Student Progressive Report"))
    elements.append(Spacer(1, 10))

    # Student Metadata Box
    user = ctx["user"]
    prog = ctx["program"]
    dept = ctx["department"]

    meta_data = [
        [
            Paragraph("<b>Student Name:</b>", bold_style), Paragraph(escape(str(user.display_name)), normal_style),
            Paragraph("<b>Reg / Roll No:</b>", bold_style), Paragraph(escape(str(student.roll_no)), normal_style)
        ],
        [
            Paragraph("<b>Programme:</b>", bold_style), Paragraph(escape(str(prog.name if prog else "N/A")), normal_style),
            Paragraph("<b>Department:</b>", bold_style), Paragraph(escape(str(dept.name if dept else "N/A")), normal_style)
        ],
        [
            Paragraph("<b>Current Semester:</b>", bold_style), Paragraph(f"Semester {student.current_semester}", normal_style),
            Paragraph("<b>Academic Standing:</b>", bold_style), Paragraph(f"<b>{ctx['overall_standing']}</b>", normal_style)
        ],
        [
            Paragraph("<b>Cumulative GPA:</b>", bold_style), Paragraph(f"<b>{ctx['cgpa']} / 4.0</b>", normal_style),
            Paragraph("<b>Credits Completed:</b>", bold_style), Paragraph(f"{ctx['earned_credits']} / {ctx['required_credits']} CR", normal_style)
        ],
    ]

    meta_table = Table(meta_data, colWidths=[90, doc.width / 2 - 90, 90, doc.width / 2 - 90])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8F9FE')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 12))

    # Semesters Tables
    for sem in ctx["semesters"]:
        sem_elements = []
        sem_elements.append(Paragraph(f"<b>{escape(sem['term_name'])}</b> (Term GPA: {sem['term_gpa']} | CGPA: {sem['cumulative_gpa']})", bold_style))
        sem_elements.append(Spacer(1, 4))

        col_widths = [55, doc.width - 55 - 42 * 5 - 60, 42, 42, 42, 42, 42, 60]
        t_data = [[
            Paragraph("Code", table_header_style),
            Paragraph("Course Unit Title", table_header_style),
            Paragraph("Credits", table_header_style),
            Paragraph("CA Mark", table_header_style),
            Paragraph("Exam", table_header_style),
            Paragraph("Total %", table_header_style),
            Paragraph("Grade", table_header_style),
            Paragraph("Outcome", table_header_style),
        ]]

        for c in sem["courses"]:
            t_data.append([
                Paragraph(escape(str(c["code"])), table_cell_style),
                Paragraph(escape(str(c["title"])), table_cell_style),
                Paragraph(str(c["credits"]), table_cell_style),
                Paragraph(str(c["ca_mark"]), table_cell_style),
                Paragraph(str(c["exam_mark"]), table_cell_style),
                Paragraph(f"{c['total_pct']}%", table_cell_style),
                Paragraph(escape(str(c["grade"])), table_cell_style),
                Paragraph(escape(str(c["outcome"])), table_cell_style),
            ])

        # Summary Row
        t_data.append([
            Paragraph("<b>SUBTOTAL:</b>", table_cell_style),
            Paragraph(f"Attempted: {sem['credits_attempted']} CR | Passed: {sem['credits_completed']} CR", table_cell_style),
            Paragraph("", table_cell_style), Paragraph("", table_cell_style), Paragraph("", table_cell_style),
            Paragraph("", table_cell_style),
            Paragraph(f"<b>GPA: {sem['term_gpa']}</b>", table_cell_style),
            Paragraph(f"<b>CGPA: {sem['cumulative_gpa']}</b>", table_cell_style),
        ])

        course_table = Table(t_data, colWidths=col_widths, repeatRows=1)
        course_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6C5CE7')),
            ('ALIGN', (2, 0), (6, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#EDF2F7')),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))

        sem_elements.append(course_table)
        sem_elements.append(Spacer(1, 10))
        elements.extend(sem_elements)

    # Signature Block (Oxford layout with optional Verification QR)
    elements.append(Spacer(1, 15))
    if verify_url:
        qr_flowable = make_qr_code_flowable(verify_url, size=46)
        sig_data = [
            [
                Paragraph("<b>Prepared By:</b><br/>Directorate of Academic Affairs<br/><br/>_____________________<br/><b>Examinations Officer</b>", normal_style),
                [
                    Paragraph("<b>Registry Verification:</b>", ParagraphStyle('SealTitle', parent=normal_style, alignment=1, fontName='Quicksand-Bold', fontSize=8)),
                    Spacer(1, 2),
                    qr_flowable,
                    Spacer(1, 2),
                    Paragraph('<font size="6.5" color="#64748b">Scan to Verify</font>', ParagraphStyle('SealSub', parent=normal_style, alignment=1)),
                ],
                Paragraph("<b>Verified &amp; Certified By:</b><br/>Dean of School / Registrar<br/><br/>_____________________<br/><b>Official University Seal</b>", normal_style),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[doc.width * 0.38, doc.width * 0.24, doc.width * 0.38])
    else:
        sig_data = [
            [Paragraph("<b>Prepared By:</b> Directorate of Academic Affairs", normal_style), Paragraph("<b>Verified &amp; Certified By:</b> Dean of School / Registrar", normal_style)],
            [Paragraph("Date Issued: " + ctx["issued_date"], normal_style), Paragraph("Official University Seal / Signature: __________________", normal_style)]
        ]
        sig_table = Table(sig_data, colWidths=[doc.width / 2] * 2)

    sig_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (1, 0), (1, 0), 'CENTER'),
        ('PADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(KeepTogether([sig_table]))

    def canvas_factory(*args, **kwargs):
        return PageNumberCanvas(*args, tracking_info=tracking_info, **kwargs)

    doc.build(elements, canvasmaker=canvas_factory)
    buf.seek(0)
    return buf.getvalue()


def export_progressive_report_excel(student, ctx):
    """
    Generates formatted .xlsx Excel workbook via openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Progressive Report"

    # Styling definitions
    font_title = Font(name="Quicksand", size=15, bold=True, color="1F2937")
    font_subtitle = Font(name="Quicksand", size=11, bold=True, color="6C5CE7")
    font_header = Font(name="Quicksand", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Quicksand", size=9.5, bold=True, color="1F2937")
    font_normal = Font(name="Quicksand", size=9.5, color="374151")

    fill_header = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
    fill_subtotal = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title
    ws.append([get_branding()["site_name"].upper()])
    ws.cell(row=1, column=1).font = font_title
    ws.append(["OFFICIAL STUDENT PROGRESSIVE REPORT"])
    ws.cell(row=2, column=1).font = font_subtitle
    ws.append([])

    # Student Metadata
    user = ctx["user"]
    prog = ctx["program"]
    dept = ctx["department"]

    meta_rows = [
        [f"Student Name: {user.display_name}", "", f"Reg / Roll No: {student.roll_no}"],
        [f"Programme: {prog.name if prog else 'N/A'}", "", f"Department: {dept.name if dept else 'N/A'}"],
        [f"Current Semester: Semester {student.current_semester}", "", f"Academic Standing: {ctx['overall_standing']}"],
        [f"Cumulative GPA: {ctx['cgpa']} / 4.0", "", f"Credits Earned: {ctx['earned_credits']} / {ctx['required_credits']}"],
    ]

    for r in meta_rows:
        ws.append(r)
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=1).font = font_bold
        ws.cell(row=row_idx, column=3).font = font_bold
        ws.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=9)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)

    ws.append([])

    headers = [
        "Code",
        "Course Title",
        "Credits",
        "CA Mark",
        "Final Exam",
        "Total Score (%)",
        "Grade",
        "GP",
        "Status"
    ]

    for sem in ctx["semesters"]:
        ws.append([f"{sem['term_name']} (Term GPA: {sem['term_gpa']} | CGPA: {sem['cumulative_gpa']})"])
        hdr_row = ws.max_row
        ws.cell(row=hdr_row, column=1).font = font_subtitle

        ws.append(headers)
        h_row_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row_idx, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for c in sem["courses"]:
            row_vals = [
                c["code"],
                c["title"],
                c["credits"],
                c["ca_mark"],
                c["exam_mark"],
                c["total_pct"],
                c["grade"],
                c["grade_point"],
                c["outcome"],
            ]
            ws.append(row_vals)
            curr_row = ws.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = font_normal
                cell.border = thin_border
                if col_idx in [3, 4, 5, 6, 7, 8]:
                    cell.alignment = Alignment(horizontal="center")

        # Subtotal Row
        subtotal_vals = [
            "SUBTOTAL",
            f"Attempted: {sem['credits_attempted']} CR | Passed: {sem['credits_completed']} CR",
            "", "", "", "", "",
            f"Term GPA: {sem['term_gpa']}",
            f"CGPA: {sem['cumulative_gpa']}"
        ]
        ws.append(subtotal_vals)
        st_row_idx = ws.max_row
        for col_idx in range(1, len(subtotal_vals) + 1):
            cell = ws.cell(row=st_row_idx, column=col_idx)
            cell.font = font_bold
            cell.fill = fill_subtotal
            cell.border = thin_border

        ws.append([])

    # Adjust Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    ws.merge_cells('A1:I1')
    ws.merge_cells('A2:I2')
    finish_worksheet(ws, header_row=10 if ctx['semesters'] else 8, filters=False)
    ws.print_title_rows = '1:7'
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_progressive_report_csv(student, ctx):
    """
    Generates CSV export representation of Student Progressive Report with UTF-8 BOM.
    """
    buf = io.StringIO()
    buf.write("\ufeff")  # UTF-8 BOM for clean spreadsheet opening
    writer = csv.writer(buf)

    writer.writerow(["Student Name", "Registration Number", "Programme", "Department", "Academic Standing", "Overall CGPA", "Term", "Course Code", "Course Title", "Credits", "CA Mark", "Final Exam", "Total Score (%)", "Letter Grade", "Grade Point", "Outcome", "Credits Attempted", "Credits Earned", "Term GPA", "Cumulative GPA"])
    for sem in ctx["semesters"]:
        for c in sem["courses"]:
            writer.writerow([ctx["user"].display_name, student.roll_no,
                ctx["program"].name if ctx["program"] else "",
                ctx["department"].name if ctx["department"] else "",
                ctx["overall_standing"], ctx["cgpa"], sem["term_name"],
                c["code"], c["title"], c["credits"], c["ca_mark"], c["exam_mark"],
                c["total_pct"], c["grade"], c["grade_point"], c["outcome"],
                sem["credits_attempted"], sem["credits_completed"], sem["term_gpa"], sem["cumulative_gpa"]])
    return buf.getvalue()
