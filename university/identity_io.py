"""
Import/export for the central user directory.

Exports respect the caller's active filters and deliberately omit passwords,
password hashes, reset tokens and any provider secret — a user export is a
directory listing, never a credential dump.
"""

import csv
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

from university.document_design import (PageNumberCanvas, ReportDocTemplate,
                                        document_styles, get_branding, letterhead)
from university.identity_models import UserType

EXPORT_HEADERS = [
    "Username", "Full Name", "User Type", "Role", "Institutional Email",
    "Contact Email", "Phone", "Student / Staff ID", "Department", "Programme",
    "Campus", "Status", "Groups", "Last Login", "Created",
]

STUDENT_IMPORT_COLUMNS = [
    ("student_id", "Student ID", False),
    ("registration_number", "Registration Number", True),
    ("first_name", "First Name", True),
    ("last_name", "Last Name", True),
    ("email", "Email", False),
    ("phone", "Phone", False),
    ("programme_code", "Programme Code", False),
    ("department_code", "Department Code", False),
    ("campus", "Campus", False),
    ("group_code", "Group Code", False),
    ("status", "Status", False),
]

STAFF_IMPORT_COLUMNS = [
    ("staff_id", "Staff ID", True),
    ("first_name", "First Name", True),
    ("last_name", "Last Name", True),
    ("email", "Email", False),
    ("phone", "Phone", False),
    ("department_code", "Department Code", False),
    ("designation", "Designation", False),
    ("role_code", "Role Code", False),
    ("group_code", "Group Code", False),
    ("campus", "Campus", False),
    ("status", "Status", False),
]

SAMPLE_ROWS = {
    UserType.STUDENT: {
        "student_id": "STU2026101", "registration_number": "BSC/001/2026",
        "first_name": "Jane", "last_name": "Wanjiru", "email": "",
        "phone": "0712345678", "programme_code": "BSC-CS", "department_code": "SCIT",
        "campus": "Main Campus", "group_code": "students", "status": "PENDING",
    },
    UserType.STAFF: {
        "staff_id": "EMP-2026-014", "first_name": "Daniel", "last_name": "Otieno",
        "email": "", "phone": "0722334455", "department_code": "SCIT",
        "designation": "Lecturer", "role_code": "lecturer", "group_code": "faculty",
        "campus": "Main Campus", "status": "PENDING",
    },
}


def import_columns_for(user_type):
    return STUDENT_IMPORT_COLUMNS if user_type == UserType.STUDENT else STAFF_IMPORT_COLUMNS


# ==============================================================================
# ROW SERIALISATION
# ==============================================================================

def user_row(user):
    """Flatten a user into the export column order."""
    account = getattr(user, "account", None)
    student = getattr(user, "student_profile", None)
    staff = getattr(user, "faculty_profile", None)
    primary_email = next((e.address for e in user.institutional_emails.all() if e.is_primary), "")

    department = ""
    if staff and staff.department:
        department = staff.department.code
    elif student and student.program and student.program.department:
        department = student.program.department.code

    return [
        user.username,
        user.get_full_name() or user.username,
        account.get_user_type_display() if account else "—",
        user.get_role_display(),
        primary_email or "—",
        user.email or "—",
        user.phone or "—",
        (getattr(student, "roll_no", "") or getattr(staff, "employee_id", "") or "—"),
        department or "—",
        (student.program.code if student and student.program else "—"),
        (account.campus if account and account.campus else "—"),
        account.effective_status if account else "—",
        ", ".join(m.group.name for m in user.group_memberships.all()) or "—",
        user.last_login.strftime("%Y-%m-%d %H:%M") if user.last_login else "Never",
        user.date_joined.strftime("%Y-%m-%d"),
    ]


# ==============================================================================
# EXPORTS
# ==============================================================================

def export_users_csv(queryset):
    buffer = io.StringIO()
    buffer.write("﻿")  # BOM so Excel picks up UTF-8
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for user in queryset:
        writer.writerow(user_row(user))
    return buffer.getvalue().encode("utf-8")


def export_users_excel(queryset, site_name="University Management System"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"

    header_fill = PatternFill("solid", fgColor="4834D4")
    title_fill = PatternFill("solid", fgColor="6C5CE7")
    zebra = PatternFill("solid", fgColor="F7F8FC")
    thin = Side(style="thin", color="D6D9E6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    span = len(EXPORT_HEADERS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    title = ws.cell(row=1, column=1, value=f"{site_name} — User Directory")
    title.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    title.fill = title_fill
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for col, header in enumerate(EXPORT_HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    for index, user in enumerate(queryset, start=3):
        for col, value in enumerate(user_row(user), start=1):
            cell = ws.cell(row=index, column=col, value=value)
            cell.font = Font(name="Arial", size=10, color="2B2B3A")
            cell.border = border
            if index % 2 == 1:
                cell.fill = zebra

    widths = [18, 26, 14, 16, 30, 28, 14, 20, 14, 14, 16, 14, 26, 18, 14]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def export_users_pdf(queryset, site_name="University Management System", logo_path=None,
                     filter_text=None):
    """Landscape directory listing using the shared branded document template."""
    buffer = io.BytesIO()
    doc = ReportDocTemplate(buffer, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
                            topMargin=36, bottomMargin=45)
    document_styles()

    th = ParagraphStyle("UserTH", fontName="Quicksand-Bold", fontSize=8, leading=10.5,
                        textColor=colors.white)
    td = ParagraphStyle("UserTD", fontName="Quicksand", fontSize=7.5, leading=10,
                        textColor=colors.HexColor("#2B2B3A"))

    branding = get_branding()
    branding.update(site_name=site_name, logo_path=logo_path or branding["logo_path"])

    story = [letterhead(doc.width, "User Directory",
                        subtitle=f"Total accounts: {queryset.count()}",
                        filter_text=filter_text, branding=branding),
             Spacer(1, 10)]

    # Credentials are intentionally absent — a directory export must never
    # become a way to exfiltrate authentication material.
    columns = ["Username", "Full Name", "Type", "Institutional Email", "ID",
               "Dept", "Status", "Last Login"]
    picks = [0, 1, 2, 4, 7, 8, 11, 13]
    data = [[Paragraph(c, th) for c in columns]]
    for user in queryset:
        row = user_row(user)
        data.append([Paragraph(str(row[i]), td) for i in picks])

    table = Table(data, repeatRows=1,
                  colWidths=[90, 130, 60, 165, 90, 55, 70, 85])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4834D4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6D9E6")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)

    doc.build(story, canvasmaker=PageNumberCanvas)
    return buffer.getvalue()


# ==============================================================================
# BULK IMPORT TEMPLATES & PARSING
# ==============================================================================

def import_template_csv(user_type):
    columns = import_columns_for(user_type)
    buffer = io.StringIO()
    buffer.write("﻿")
    writer = csv.writer(buffer)
    writer.writerow([label for _key, label, _req in columns])
    sample = SAMPLE_ROWS[user_type]
    writer.writerow([sample.get(key, "") for key, _label, _req in columns])
    return buffer.getvalue().encode("utf-8")


def import_template_excel(user_type, site_name="University Management System"):
    columns = import_columns_for(user_type)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import Template"

    header_fill = PatternFill("solid", fgColor="4834D4")
    required_fill = PatternFill("solid", fgColor="E84393")
    for col, (_key, label, required) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=label + (" *" if required else ""))
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = required_fill if required else header_fill
        ws.column_dimensions[get_column_letter(col)].width = max(16, len(label) + 6)

    sample = SAMPLE_ROWS[user_type]
    for col, (key, _label, _req) in enumerate(columns, start=1):
        ws.cell(row=2, column=col, value=sample.get(key, ""))

    notes = ws.cell(row=4, column=1,
                    value="Columns marked * are required. Usernames, institutional emails "
                          "and passwords are generated by the system — never enter them here.")
    notes.font = Font(name="Arial", size=10, italic=True, color="718096")
    ws.freeze_panes = "A2"

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def parse_import_file(uploaded_file, user_type):
    """
    Read a CSV or Excel upload into normalised row dicts.

    Returns ``(rows, error)``. Header matching is case-insensitive and tolerant
    of the '*' suffix used to mark required columns in the template.
    """
    columns = import_columns_for(user_type)
    label_to_key = {label.lower(): key for key, label, _req in columns}
    label_to_key.update({key.lower(): key for key, _label, _req in columns})

    name = (uploaded_file.name or "").lower()
    raw_rows = []

    try:
        if name.endswith((".xlsx", ".xlsm", ".xls")):
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            ws = wb.active
            iterator = ws.iter_rows(values_only=True)
            headers = [str(h or "").strip().rstrip("*").strip().lower()
                       for h in next(iterator, [])]
            for values in iterator:
                raw_rows.append(dict(zip(headers, values)))
        else:
            content = uploaded_file.read()
            if isinstance(content, bytes):
                content = content.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(content))
            for row in reader:
                raw_rows.append({str(k or "").strip().rstrip("*").strip().lower(): v
                                 for k, v in row.items()})
    except Exception as exc:
        return [], f"The file could not be read: {exc}"

    rows = []
    for index, raw in enumerate(raw_rows, start=2):
        if not any(str(v or "").strip() for v in raw.values()):
            continue
        record = {"_row": index}
        for header, value in raw.items():
            key = label_to_key.get(str(header or "").strip().lower())
            if key:
                record[key] = str(value).strip() if value is not None else ""
        rows.append(record)

    if not rows:
        return [], "No data rows were found in the uploaded file."
    return rows, None
