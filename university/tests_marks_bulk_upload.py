import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

import openpyxl

from accounts.models import FacultyProfile, Role, StudentProfile
from university.models import (
    AcademicTerm, Course, Department, Exam, Program, Result, Assignment, Submission
)
from university import marks_io
from university import examination_services as workflow

User = get_user_model()


class MarksBulkUploadTestCase(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Computer Science", code="CS")
        self.prog = Program.objects.create(name="BSc CS", code="BCS", department=self.dept)

        # Faculty
        self.fac_user = User.objects.create_user(
            username="prof.smith", password="password123", role=Role.FACULTY,
            first_name="Prof", last_name="Smith"
        )
        self.fac = FacultyProfile.objects.create(user=self.fac_user, employee_id="FAC001", department=self.dept)

        # Admin
        self.admin_user = User.objects.create_user(
            username="admin.test", password="password123", role=Role.ADMIN,
            is_staff=True, is_superuser=True
        )

        # Course & Term
        self.course = Course.objects.create(
            code="CS101", title="Intro to CS", department=self.dept, faculty=self.fac
        )
        self.term = AcademicTerm.objects.create(
            name="Fall 2026", start_date="2026-09-01", end_date="2026-12-31"
        )

        # Students
        self.students = []
        for i in range(1, 4):
            u = User.objects.create_user(
                username=f"student{i}", password="password123", role=Role.STUDENT,
                first_name=f"Student{i}", last_name="Test"
            )
            sp = StudentProfile.objects.create(user=u, roll_no=f"STU00{i}", program=self.prog)
            self.students.append(sp)

        # Exam sitting
        self.exam = Exam.objects.create(
            name="Main Sitting", course=self.course, term=self.term,
            cat_max_marks=Decimal("30.00"), exam_max_marks=Decimal("70.00"), max_marks=100,
            status=Exam.Status.MARKING
        )
        for i, s in enumerate(self.students, 1):
            Result.objects.create(exam=self.exam, student=s, seat_number=i, attendance="PENDING")

        # Assignment
        self.assignment = Assignment.objects.create(
            course=self.course, title="Homework 1", max_marks=50, due_date="2026-10-31"
        )
        for s in self.students:
            Submission.objects.create(assignment=self.assignment, student=s, content="My homework")

    def test_generate_exam_marks_template_csv(self):
        csv_bytes = marks_io.generate_exam_marks_template(self.exam, fmt="csv")
        self.assertIsInstance(csv_bytes, bytes)
        content = csv_bytes.decode("utf-8")
        self.assertIn("roll_no", content)
        self.assertIn("STU001", content)
        self.assertIn("STU002", content)
        self.assertIn("STU003", content)

    def test_generate_exam_marks_template_excel(self):
        excel_bytes = marks_io.generate_exam_marks_template(self.exam, fmt="excel")
        self.assertIsInstance(excel_bytes, bytes)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        ws = wb.active
        self.assertEqual(ws.title, "Marks Entry")
        # Header should contain roll numbers
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
        self.assertIn("STU001", values)
        self.assertIn("STU002", values)
        self.assertIn("STU003", values)

    def test_parse_exam_marks_csv_valid(self):
        csv_content = (
            "Roll Number,Student Name,Attendance,CAT Marks,Exam Marks,Remarks\n"
            "STU001,Student1 Test,PRESENT,25.5,58.0,Good work\n"
            "STU002,Student2 Test,PRESENT,18.0,42.5,Fair\n"
            "STU003,Student3 Test,ABSENT,,,\n"
        ).encode("utf-8")

        upload = SimpleUploadedFile("marks.csv", csv_content, content_type="text/csv")
        parsed = marks_io.parse_exam_marks_file(upload, self.exam)

        self.assertEqual(parsed["valid_count"], 3)
        self.assertEqual(parsed["error_count"], 0)
        self.assertEqual(len(parsed["entries"]), 3)

        res1 = Result.objects.get(exam=self.exam, student__roll_no="STU001")
        entry1 = parsed["entries"][str(res1.pk)]
        self.assertEqual(entry1["attendance"], "PRESENT")
        self.assertEqual(entry1["cat_marks"], "25.50")
        self.assertEqual(entry1["exam_marks"], "58.00")
        self.assertEqual(entry1["marks"], "83.50")

    def test_parse_exam_marks_excel_valid(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Seat", "Roll Number", "Name", "Attendance", "CAT Marks", "Exam Marks", "Remarks"])
        ws.append([1, "STU001", "Student1", "PRESENT", 28, 65, "Excellent"])
        ws.append([2, "STU002", "Student2", "PRESENT", 15, 35, "Pass"])
        ws.append([3, "STU003", "Student3", "ABSENT", None, None, "Sick"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        upload = SimpleUploadedFile("marks.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        parsed = marks_io.parse_exam_marks_file(upload, self.exam)

        self.assertEqual(parsed["valid_count"], 3)
        self.assertEqual(parsed["error_count"], 0)

    def test_parse_exam_marks_validation_errors(self):
        csv_content = (
            "Roll Number,Attendance,CAT Marks,Exam Marks\n"
            "STU001,PRESENT,35.0,50.0\n"  # CAT > 30.0 (error)
            "STU002,ABSENT,10.0,40.0\n"   # Absent candidate with marks (error)
            "STU999,PRESENT,20.0,40.0\n"   # Unregistered candidate (error)
        ).encode("utf-8")

        upload = SimpleUploadedFile("marks.csv", csv_content, content_type="text/csv")
        parsed = marks_io.parse_exam_marks_file(upload, self.exam)

        self.assertEqual(parsed["valid_count"], 0)
        self.assertEqual(parsed["error_count"], 3)

    def test_assignment_marks_template_and_parse(self):
        # Generate template
        excel_bytes = marks_io.generate_assignment_marks_template(self.assignment, fmt="excel")
        self.assertIsInstance(excel_bytes, bytes)

        # Parse valid upload
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Roll Number", "Name", "Marks", "Max Marks", "Feedback"])
        ws.append(["STU001", "Student1", 45, 50, "Great!"])
        ws.append(["STU002", "Student2", 38, 50, "Well done"])
        ws.append(["STU003", "Student3", 25, 50, "Needs improvement"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        upload = SimpleUploadedFile("assignment.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        parsed = marks_io.parse_assignment_marks_file(upload, self.assignment)

        self.assertEqual(parsed["valid_count"], 3)
        self.assertEqual(parsed["error_count"], 0)
        self.assertEqual(len(parsed["updated_subs"]), 3)

    def test_views_exam_marks_excel_template_download(self):
        self.client.force_login(self.fac_user)
        url = reverse("examinations:marks", args=[self.exam.pk]) + "?format=excel&template=1"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

    def test_views_exam_marks_bulk_upload_post(self):
        self.client.force_login(self.fac_user)
        csv_content = (
            "Roll Number,Student Name,Attendance,CAT Marks,Exam Marks,Remarks\n"
            "STU001,Student1 Test,PRESENT,22.0,60.0,Passed\n"
            "STU002,Student2 Test,PRESENT,20.0,50.0,Passed\n"
            "STU003,Student3 Test,PRESENT,28.0,68.0,Passed\n"
        ).encode("utf-8")

        upload = SimpleUploadedFile("marks.csv", csv_content, content_type="text/csv")
        url = reverse("examinations:marks", args=[self.exam.pk])
        response = self.client.post(url, {
            "marks_file": upload,
            "revision": self.exam.revision,
        })
        self.assertEqual(response.status_code, 302)

        # Check that results were updated in the database
        res1 = Result.objects.get(exam=self.exam, student__roll_no="STU001")
        self.assertEqual(res1.marks_obtained, Decimal("82.00"))
        self.assertEqual(res1.grade, "A")
