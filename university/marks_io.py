import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from university.models import Exam, Result, Submission


# ==============================================================================
# COLUMN MAPPING HELPERS
# ==============================================================================

def _normalize_header(header):
    if not header:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(header).strip().lower())


HEADER_MAP_EXAM = {
    "roll_no": ["rollno", "rollnumber", "roll", "regno", "registrationno", "registrationnumber", "studentid"],
    "student_name": ["studentname", "name", "candidate", "candidatename", "fullname"],
    "attendance": ["attendance", "attendancestatus", "status", "att"],
    "cat_marks": ["catmarks", "cat", "catmark", "coursework", "ca", "catmark30"],
    "exam_marks": ["exammarks", "exam", "exammark", "finalexam", "examination", "exammark70"],
    "total_marks": ["totalmarks", "total", "marks", "marksobtained", "finalmark", "score"],
    "remarks": ["remarks", "remark", "comments", "comment", "notes"],
}

HEADER_MAP_ASSIGNMENT = {
    "roll_no": ["rollno", "rollnumber", "roll", "regno", "registrationno", "registrationnumber", "studentid"],
    "student_name": ["studentname", "name", "candidate", "fullname"],
    "marks": ["marks", "mark", "score", "grade", "marksobtained"],
    "feedback": ["feedback", "comments", "remarks", "comment"],
}


def _match_columns(raw_headers, header_map):
    mapping = {}
    for idx, raw in enumerate(raw_headers):
        clean = _normalize_header(raw)
        if not clean:
            continue
        for key, aliases in header_map.items():
            if clean in aliases and key not in mapping:
                mapping[key] = idx
                break
    return mapping


def _parse_decimal(val):
    if val is None or val == "":
        return None
    try:
        s = str(val).strip()
        if not s or s.upper() in ("N/A", "NONE", "—", "-"):
            return None
        return round(Decimal(s), 2)
    except (InvalidOperation, ValueError, TypeError):
        return "INVALID"


# ==============================================================================
# 1. EXAMINATION MARKS TEMPLATE GENERATOR
# ==============================================================================

def generate_exam_marks_template(exam, fmt="excel"):
    """
    Generate an Excel (.xlsx) or CSV (.csv) template pre-populated with
    registered candidates for this examination sitting.
    """
    rows = list(
        exam.results.select_related("student__user", "exam").order_by(
            "seat_number", "student__roll_no"
        )
    )

    if fmt == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")  # UTF-8 BOM
        writer = csv.writer(buf)
        writer.writerow(["roll_no", "student_name", "attendance", "cat_marks", "exam_marks", "remarks"])
        for r in rows:
            writer.writerow([
                r.student.roll_no,
                r.student.user.display_name,
                r.attendance if r.attendance else "PENDING",
                f"{r.cat_marks:.2f}" if r.cat_marks is not None else "",
                f"{r.exam_marks:.2f}" if r.exam_marks is not None else "",
                r.remarks or "",
            ])
        return buf.getvalue().encode("utf-8")

    # Generate Excel (.xlsx)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks Entry"

    # Styling Palette
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Title & Metadata
    ws["A1"] = f"{exam.course.code} — {exam.course.title} · Marks Entry Template"
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    ws.merge_cells("A1:G1")

    cat_max = float(exam.cat_max_marks) if exam.cat_max_marks is not None else 30.0
    exam_max = float(exam.exam_max_marks) if exam.exam_max_marks is not None else 70.0
    total_max = exam.max_marks or 100

    ws["A2"] = (
        f"Term: {exam.term.name if exam.term else 'Current'} | "
        f"Exam: {exam.name} | "
        f"CAT Max: {cat_max:.0f} | "
        f"Final Exam Max: {exam_max:.0f} | "
        f"Total Max Marks: {total_max} | "
        f"Pass Mark: {exam.pass_mark:.0f}%"
    )
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="475569")
    ws.merge_cells("A2:G2")

    ws["A3"] = "Instructions: Enter 'PRESENT' or 'ABSENT' under Attendance. Fill in CAT Marks and Exam Marks. Total and CUE Letter Grades are computed automatically."
    ws["A3"].font = Font(name="Arial", size=8.5, color="64748B")
    ws.merge_cells("A3:G3")

    # Column Headers (Row 5)
    headers = [
        ("Seat No", 10),
        ("Roll Number", 18),
        ("Student Name", 28),
        ("Attendance", 16),
        (f"CAT Marks (Max {cat_max:.0f})", 20),
        (f"Exam Marks (Max {exam_max:.0f})", 22),
        ("Remarks", 26),
    ]

    header_row = 5
    for col_num, (h_title, width) in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_num, value=h_title)
        c.font = Font(name="Arial", size=10, bold=True, color="0F172A")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_num in (1, 4, 5, 6) else "left", vertical="center")
        c.border = thin_border
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = max(width, len(h_title) + 4)

    # Candidate Rows
    curr_row = 6
    for r in rows:
        ws.cell(row=curr_row, column=1, value=r.seat_number or "").alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=2, value=r.student.roll_no).alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=3, value=r.student.user.display_name).alignment = Alignment(horizontal="left")
        
        att_val = r.attendance if r.attendance else "PENDING"
        ws.cell(row=curr_row, column=4, value=att_val).alignment = Alignment(horizontal="center")
        
        cat_c = ws.cell(row=curr_row, column=5, value=float(r.cat_marks) if r.cat_marks is not None else "")
        cat_c.alignment = Alignment(horizontal="right")
        cat_c.number_format = "0.00"
        
        exam_c = ws.cell(row=curr_row, column=6, value=float(r.exam_marks) if r.exam_marks is not None else "")
        exam_c.alignment = Alignment(horizontal="right")
        exam_c.number_format = "0.00"

        ws.cell(row=curr_row, column=7, value=r.remarks or "").alignment = Alignment(horizontal="left")

        for c_idx in range(1, 8):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.font = Font(name="Arial", size=9.5)
            cell.border = thin_border
            if curr_row % 2 == 1:
                cell.fill = light_blue_fill

        curr_row += 1

    # Freeze panes below headers
    ws.freeze_panes = "A6"
    ws.views.sheetView[0].showGridLines = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ==============================================================================
# 2. EXAMINATION MARKS FILE PARSER
# ==============================================================================

def parse_exam_marks_file(upload_file, exam):
    """
    Parse an uploaded CSV or Excel file containing examination marks.
    Returns:
      {
        "valid_count": int,
        "error_count": int,
        "duplicate_count": int,
        "total_count": int,
        "entries": {result_id_str: record_dict},
        "items": [row_detail_dict, ...],
        "general_errors": [str, ...]
      }
    """
    # Verify file extension & size
    filename = getattr(upload_file, "name", "marks.csv").lower()
    if upload_file.size > 5 * 1024 * 1024:
        raise ValidationError("File size exceeds maximum allowed 5 MB.")

    is_excel = filename.endswith(".xlsx") or filename.endswith(".xls")
    is_csv = filename.endswith(".csv")

    if not (is_excel or is_csv):
        raise ValidationError("Unsupported file format. Please upload a CSV (.csv) or Excel (.xlsx / .xls) spreadsheet.")

    # Read records from file
    raw_rows = []
    if is_excel:
        try:
            wb = openpyxl.load_workbook(upload_file, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    raw_rows.append([str(c).strip() if c is not None else "" for c in row])
        except Exception as e:
            raise ValidationError(f"Could not read Excel file: {str(e)}")
    else:
        content = upload_file.read()
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                raise ValidationError("CSV encoding not recognized. Please save file as UTF-8.")
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            if any(cell.strip() for cell in row):
                raw_rows.append([cell.strip() for cell in row])

    if not raw_rows:
        raise ValidationError("The uploaded spreadsheet is empty.")

    # Locate Header Row
    header_idx = -1
    col_map = {}
    for idx, row in enumerate(raw_rows[:15]):  # Check first 15 rows for header
        mapped = _match_columns(row, HEADER_MAP_EXAM)
        if "roll_no" in mapped:
            header_idx = idx
            col_map = mapped
            break

    if header_idx == -1 or "roll_no" not in col_map:
        raise ValidationError(
            "Could not identify required 'Roll Number' column in header. "
            "Please ensure your spreadsheet has a header row with Roll Number, Attendance, CAT Marks, and Exam Marks."
        )

    # Prepare Candidate Lookup
    candidates = {
        r.student.roll_no.strip().upper(): r
        for r in exam.results.select_related("student__user", "exam")
    }

    cat_max = exam.cat_max_marks if exam.cat_max_marks is not None else Decimal(30)
    exam_max = exam.exam_max_marks if exam.exam_max_marks is not None else Decimal(70)
    total_max = exam.max_marks or 100

    data_rows = raw_rows[header_idx + 1:]
    items = []
    entries = {}
    seen_rolls = set()
    error_count = 0
    duplicate_count = 0
    valid_count = 0

    for row_num, row in enumerate(data_rows, start=header_idx + 2):
        # Extract fields
        roll_idx = col_map["roll_no"]
        roll_val = row[roll_idx].strip() if roll_idx < len(row) else ""
        if not roll_val:
            continue

        clean_roll = roll_val.upper()
        row_errors = []

        # Candidate check
        result_obj = candidates.get(clean_roll)
        if not result_obj:
            row_errors.append(f"Student roll number '{roll_val}' is not registered for this examination.")

        # Duplicate check in spreadsheet
        if clean_roll in seen_rolls:
            row_errors.append(f"Duplicate candidate '{roll_val}' in spreadsheet.")
            duplicate_count += 1
            status = "duplicate"
        else:
            status = "valid"
        seen_rolls.add(clean_roll)

        # Attendance check
        att_idx = col_map.get("attendance")
        raw_att = (row[att_idx].strip().upper() if att_idx is not None and att_idx < len(row) else "")
        attendance = "PENDING"
        if raw_att in ("PRESENT", "P", "1", "YES", "TRUE"):
            attendance = "PRESENT"
        elif raw_att in ("ABSENT", "A", "0", "NO", "FALSE"):
            attendance = "ABSENT"
        elif raw_att in ("PENDING", ""):
            attendance = "PENDING"
        else:
            row_errors.append(f"Invalid attendance '{raw_att}'. Must be PRESENT or ABSENT.")

        # CAT marks
        cat_idx = col_map.get("cat_marks")
        raw_cat = row[cat_idx] if cat_idx is not None and cat_idx < len(row) else ""
        cat_parsed = _parse_decimal(raw_cat)
        if cat_parsed == "INVALID":
            row_errors.append(f"CAT mark '{raw_cat}' is not a valid number.")
            cat_parsed = None
        elif cat_parsed is not None:
            if cat_parsed < 0 or cat_parsed > cat_max:
                row_errors.append(f"CAT mark ({cat_parsed}) exceeds allowed range 0 – {cat_max:.0f}.")

        # Exam marks
        exam_idx = col_map.get("exam_marks")
        raw_exam = row[exam_idx] if exam_idx is not None and exam_idx < len(row) else ""
        exam_parsed = _parse_decimal(raw_exam)
        if exam_parsed == "INVALID":
            row_errors.append(f"Exam mark '{raw_exam}' is not a valid number.")
            exam_parsed = None
        elif exam_parsed is not None:
            if exam_parsed < 0 or exam_parsed > exam_max:
                row_errors.append(f"Exam mark ({exam_parsed}) exceeds allowed range 0 – {exam_max:.0f}.")

        # Total marks (if provided directly)
        total_idx = col_map.get("total_marks")
        raw_total = row[total_idx] if total_idx is not None and total_idx < len(row) else ""
        total_parsed = _parse_decimal(raw_total)
        if total_parsed == "INVALID":
            total_parsed = None

        # Auto-compute total if CAT and Exam are given
        if cat_parsed is not None or exam_parsed is not None:
            calc_total = (cat_parsed or Decimal(0)) + (exam_parsed or Decimal(0))
        elif total_parsed is not None:
            calc_total = total_parsed
        else:
            calc_total = None

        # Absent logic
        if attendance == "ABSENT" and (cat_parsed is not None or exam_parsed is not None or total_parsed is not None):
            row_errors.append("Absent candidate cannot have CAT or Exam marks recorded.")

        # If marks provided but attendance is PENDING, auto-set to PRESENT
        if attendance == "PENDING" and (cat_parsed is not None or exam_parsed is not None or total_parsed is not None):
            attendance = "PRESENT"

        # Remarks
        rem_idx = col_map.get("remarks")
        remarks_val = row[rem_idx].strip() if rem_idx is not None and rem_idx < len(row) else ""

        if row_errors:
            status = "error" if status != "duplicate" else "duplicate"
            error_count += 1
        else:
            valid_count += 1
            if result_obj:
                entries[str(result_obj.pk)] = {
                    "attendance": attendance,
                    "cat_marks": str(cat_parsed) if cat_parsed is not None else "",
                    "exam_marks": str(exam_parsed) if exam_parsed is not None else "",
                    "marks": str(calc_total) if calc_total is not None else "",
                    "remarks": remarks_val,
                }

        items.append({
            "row_num": row_num,
            "roll_no": roll_val,
            "student_name": result_obj.student.user.display_name if result_obj else "—",
            "attendance": attendance,
            "cat_marks": cat_parsed,
            "exam_marks": exam_parsed,
            "total_marks": calc_total,
            "remarks": remarks_val,
            "status": status,
            "errors": row_errors,
        })

    return {
        "valid_count": valid_count,
        "error_count": error_count,
        "duplicate_count": duplicate_count,
        "total_count": len(items),
        "entries": entries,
        "items": items,
        "general_errors": [],
    }


# ==============================================================================
# 3. ASSIGNMENT MARKS TEMPLATE & PARSER
# ==============================================================================

def generate_assignment_marks_template(assignment, fmt="excel"):
    """
    Generate an Excel or CSV template pre-populated with student submissions
    for coursework grading.
    """
    subs = list(
        assignment.submissions.select_related("student__user").order_by("student__roll_no")
    )

    if fmt == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")
        writer = csv.writer(buf)
        writer.writerow(["roll_no", "student_name", "marks", "max_marks", "feedback"])
        for s in subs:
            writer.writerow([
                s.student.roll_no,
                s.student.user.display_name,
                s.marks if s.marks is not None else "",
                assignment.max_marks,
                s.feedback or "",
            ])
        return buf.getvalue().encode("utf-8")

    # Excel (.xlsx)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assignment Marks"

    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws["A1"] = f"{assignment.course.code} — {assignment.title} · Marks Template"
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    ws.merge_cells("A1:E1")

    ws["A2"] = f"Course: {assignment.course.title} | Max Marks: {assignment.max_marks} | Total Submissions: {len(subs)}"
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="475569")
    ws.merge_cells("A2:E2")

    headers = [
        ("Roll Number", 18),
        ("Student Name", 28),
        (f"Marks (Max {assignment.max_marks})", 22),
        ("Max Marks", 12),
        ("Feedback", 32),
    ]

    header_row = 4
    for col_num, (h_title, width) in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_num, value=h_title)
        c.font = Font(name="Arial", size=10, bold=True, color="0F172A")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_num in (3, 4) else "left", vertical="center")
        c.border = thin_border
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = max(width, len(h_title) + 4)

    curr_row = 5
    for s in subs:
        ws.cell(row=curr_row, column=1, value=s.student.roll_no).alignment = Alignment(horizontal="left")
        ws.cell(row=curr_row, column=2, value=s.student.user.display_name).alignment = Alignment(horizontal="left")
        
        m_c = ws.cell(row=curr_row, column=3, value=s.marks if s.marks is not None else "")
        m_c.alignment = Alignment(horizontal="right")
        
        max_c = ws.cell(row=curr_row, column=4, value=assignment.max_marks)
        max_c.alignment = Alignment(horizontal="center")
        
        ws.cell(row=curr_row, column=5, value=s.feedback or "").alignment = Alignment(horizontal="left")

        for c_idx in range(1, 6):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.font = Font(name="Arial", size=9.5)
            cell.border = thin_border
            if curr_row % 2 == 1:
                cell.fill = light_blue_fill

        curr_row += 1

    ws.freeze_panes = "A5"
    ws.views.sheetView[0].showGridLines = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def parse_assignment_marks_file(upload_file, assignment):
    """
    Parse an uploaded spreadsheet for coursework assignment grading.
    """
    filename = getattr(upload_file, "name", "assignment_marks.csv").lower()
    if upload_file.size > 5 * 1024 * 1024:
        raise ValidationError("File size exceeds maximum allowed 5 MB.")

    is_excel = filename.endswith(".xlsx") or filename.endswith(".xls")
    is_csv = filename.endswith(".csv")

    if not (is_excel or is_csv):
        raise ValidationError("Unsupported file format. Please upload CSV or Excel spreadsheet.")

    raw_rows = []
    if is_excel:
        try:
            wb = openpyxl.load_workbook(upload_file, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    raw_rows.append([str(c).strip() if c is not None else "" for c in row])
        except Exception as e:
            raise ValidationError(f"Could not read Excel file: {str(e)}")
    else:
        content = upload_file.read()
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            if any(cell.strip() for cell in row):
                raw_rows.append([cell.strip() for cell in row])

    if not raw_rows:
        raise ValidationError("The uploaded spreadsheet is empty.")

    header_idx = -1
    col_map = {}
    for idx, row in enumerate(raw_rows[:15]):
        mapped = _match_columns(row, HEADER_MAP_ASSIGNMENT)
        if "roll_no" in mapped and "marks" in mapped:
            header_idx = idx
            col_map = mapped
            break

    if header_idx == -1 or "roll_no" not in col_map or "marks" not in col_map:
        raise ValidationError(
            "Could not identify required 'Roll Number' and 'Marks' columns in header."
        )

    subs = {
        s.student.roll_no.strip().upper(): s
        for s in assignment.submissions.select_related("student__user")
    }

    items = []
    updated_subs = []
    error_count = 0
    valid_count = 0

    for row_num, row in enumerate(raw_rows[header_idx + 1:], start=header_idx + 2):
        roll_idx = col_map["roll_no"]
        roll_val = row[roll_idx].strip() if roll_idx < len(row) else ""
        if not roll_val:
            continue

        clean_roll = roll_val.upper()
        sub_obj = subs.get(clean_roll)
        row_errors = []

        if not sub_obj:
            row_errors.append(f"Student roll '{roll_val}' not found in assignment submissions.")

        marks_idx = col_map["marks"]
        raw_marks = row[marks_idx] if marks_idx < len(row) else ""
        parsed_marks = _parse_decimal(raw_marks)
        if parsed_marks == "INVALID":
            row_errors.append(f"Marks '{raw_marks}' is not a valid number.")
            parsed_marks = None
        elif parsed_marks is not None:
            if parsed_marks < 0 or parsed_marks > assignment.max_marks:
                row_errors.append(f"Marks ({parsed_marks}) must be between 0 and {assignment.max_marks}.")

        fb_idx = col_map.get("feedback")
        feedback = row[fb_idx].strip() if fb_idx is not None and fb_idx < len(row) else ""

        if row_errors:
            error_count += 1
            status = "error"
        else:
            valid_count += 1
            status = "valid"
            if sub_obj and parsed_marks is not None:
                sub_obj.marks = int(round(parsed_marks))
                sub_obj.feedback = feedback
                sub_obj.status = Submission.GRADED
                updated_subs.append(sub_obj)

        items.append({
            "row_num": row_num,
            "roll_no": roll_val,
            "student_name": sub_obj.student.user.display_name if sub_obj else "—",
            "marks": parsed_marks,
            "feedback": feedback,
            "status": status,
            "errors": row_errors,
        })

    return {
        "valid_count": valid_count,
        "error_count": error_count,
        "total_count": len(items),
        "updated_subs": updated_subs,
        "items": items,
    }
