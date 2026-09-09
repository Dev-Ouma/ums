"""Shared visual design system for every generated PDF/Excel/CSV document in
the system — one brand identity, one letterhead structure, one table style,
reused by every *_io.py exporter instead of each hand-rolling its own.

Two palettes are provided:
- The live system brand palette (matches static/css/ums.css --primary) for
  ordinary reports: course/student/faculty/fee/timetable lists, progressive
  reports, exam cards, nominal rolls.
- A muted "formal" palette for certified academic documents (transcripts)
  which must stay deliberately low-color and authoritative.
"""
import os
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from xml.sax.saxutils import escape

from django.conf import settings

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage, Paragraph, Table, TableStyle, SimpleDocTemplate


@lru_cache(maxsize=1)
def document_fonts():
    """Embed the same locally hosted font used by the web interface."""
    root = os.path.join(settings.BASE_DIR, 'static', 'fonts', 'quicksand')
    for name in ('Quicksand', 'Quicksand-Bold', 'Quicksand-Medium', 'Quicksand-SemiBold'):
        font_path = os.path.join(root, name + '.ttf')
        if os.path.exists(font_path):
            pdfmetrics.registerFont(TTFont(name, font_path))
    pdfmetrics.registerFontFamily('Quicksand', normal='Quicksand', bold='Quicksand-Bold',
                                  italic='Quicksand', boldItalic='Quicksand-Bold')
    return 'Quicksand', 'Quicksand-Bold'


def document_styles():
    from reportlab.lib.styles import getSampleStyleSheet
    regular, bold = document_fonts()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        if hasattr(style, 'fontName'):
            style.fontName = bold if 'Bold' in style.fontName else regular
        if hasattr(style, 'textColor') and style.textColor == colors.HexColor('#1E293B'):
            style.textColor = colors.HexColor(INK)
    return styles


class ReportDocTemplate(SimpleDocTemplate):
    """A4 document defaults shared by existing exporters without changing data."""
    def __init__(self, *args, kwargs_extra=None, **kwargs):
        document_fonts()
        kwargs.setdefault('pagesize', A4)
        kwargs['bottomMargin'] = max(kwargs.get('bottomMargin', 48), 48)
        kwargs.setdefault('author', get_branding()['site_name'])
        super().__init__(*args, **kwargs)

    def build(self, flowables, *args, **kwargs):
        kwargs.setdefault('canvasmaker', PageNumberCanvas)
        return super().build(flowables, *args, **kwargs)

# ==============================================================================
# BRAND PALETTE — matches the live site's --primary / --primary-dark / --accent
# ==============================================================================
PRIMARY = "#6C5CE7"
PRIMARY_DARK = "#4834D4"
ACCENT = "#E84393"
INK = "#0F172A"
MUTED = "#334155"
BORDER = "#CBD5E1"
ZEBRA = "#F8F9FE"
HEADER_FILL = PRIMARY_DARK

# ==============================================================================
# FORMAL PALETTE — certified academic documents only (transcripts)
# ==============================================================================
FORMAL_INK = "#0F172A"
FORMAL_MUTED = "#334155"
FORMAL_RULE = "#1E293B"
FORMAL_FAINT_RULE = "#CBD5E1"
FORMAL_HEAD_BG = "#F1F5F9"


def get_branding():
    """Live university name / logo / contact info from CMS Site Settings,
    with safe fallbacks so every report renders even before Setup is configured."""
    site_name = "University Management System"
    contact_address = contact_email = contact_phone = ""
    logo_path = os.path.join(settings.BASE_DIR, "static", "img", "ums-logo.png")

    try:
        from cms.models import SiteSettings
        site = SiteSettings.objects.first()
        if site:
            site_name = site.site_name or site_name
            contact_address = site.contact_address or ""
            contact_email = site.contact_email or ""
            contact_phone = site.contact_phone or ""
            if site.logo:
                try:
                    if site.logo.path and os.path.isfile(site.logo.path):
                        logo_path = site.logo.path
                except (ValueError, NotImplementedError):
                    pass
    except Exception:
        pass

    if not logo_path or not os.path.isfile(logo_path):
        from university.document_design import _generate_fallback_logo
        logo_path = _generate_fallback_logo()

    return {
        "site_name": site_name,
        "contact_address": contact_address,
        "contact_email": contact_email,
        "contact_phone": contact_phone,
        "contact_line": " · ".join(filter(None, [contact_address, contact_email, contact_phone])),
        "logo_path": logo_path if (logo_path and os.path.isfile(logo_path)) else None,
    }


def _generate_fallback_logo():
    logo_dir = os.path.join(settings.BASE_DIR, "static", "img")
    os.makedirs(logo_dir, exist_ok=True)
    logo_file = os.path.join(logo_dir, "ums-logo.png")
    if os.path.isfile(logo_file) and os.path.getsize(logo_file) > 100:
        return logo_file
    try:
        from PIL import Image, ImageDraw
        img = Image.new("RGBA", (180, 180), (108, 92, 231, 255))
        draw = ImageDraw.Draw(img)
        draw.polygon([(90, 25), (150, 60), (150, 120), (90, 155), (30, 120), (30, 60)], fill=(72, 52, 212, 255))
        draw.rectangle([(65, 65), (115, 115)], fill=(255, 255, 255, 255))
        draw.polygon([(90, 50), (130, 75), (50, 75)], fill=(232, 67, 147, 255))
        img.save(logo_file, "PNG")
        return logo_file
    except Exception:
        return None


# ==============================================================================
# REPORTLAB HELPERS
# ==============================================================================
class PageNumberCanvas(pdfcanvas.Canvas):
    """Two-pass canvas that writes dynamic page counts ("Page X of Y"),
    continuation headers on continuation pages, optional watermarks, and audit tracking."""
    def __init__(self, *args, running_header=None, footer_left=None, watermark=None, tracking_info=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self._running_header = running_header
        self._footer_left = footer_left or get_branding()['site_name']
        self._watermark = watermark
        self._tracking_info = tracking_info

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_footer(num_pages)
            if self._pageNumber > 1 and self._running_header:
                self._draw_running_header()
            if self._watermark:
                self._draw_watermark()
            pdfcanvas.Canvas.showPage(self)
        pdfcanvas.Canvas.save(self)

    def _draw_running_header(self):
        width, height = self._pagesize
        self.saveState()
        font_name = document_fonts()[0]
        self.setFont(font_name, 7.5)
        self.setFillColor(colors.HexColor(MUTED))
        self.drawString(36, height - 24, self._running_header)
        self.setStrokeColor(colors.HexColor(BORDER))
        self.setLineWidth(0.4)
        self.line(36, height - 28, width - 36, height - 28)
        self.restoreState()

    def _draw_footer(self, page_count):
        width = self._pagesize[0]
        self.saveState()
        font_name = document_fonts()[0]
        self.setFont(font_name, 7.5)
        self.setFillColor(colors.HexColor(MUTED))
        footer = self._footer_left
        while pdfmetrics.stringWidth(footer, font_name, 7.5) > width - 165:
            footer = footer[:-2]
        self.drawString(36, 22, footer)
        self.drawRightString(width - 36, 22, f"Page {self._pageNumber} of {page_count}")
        self.setStrokeColor(colors.HexColor(BORDER))
        self.setLineWidth(0.5)
        self.line(36, 32, width - 36, 32)
        if self._tracking_info:
            self.setFont(font_name, 6.5)
            self.setFillColor(colors.HexColor("#64748B"))
            self.drawString(36, 12, f"Security Audit: {self._tracking_info}")
        self.restoreState()

    def _draw_watermark(self):
        width, height = self._pagesize
        self.saveState()
        self.setFillColor(colors.HexColor("#94A3B8"))
        try:
            self.setFillAlpha(0.06)
        except Exception:
            pass
        self.setFont("Helvetica-Bold", 54)
        self.translate(width / 2, height / 2)
        self.rotate(36)
        self.drawCentredString(0, 0, self._watermark)
        self.restoreState()


def make_qr_code_flowable(data, size=52):
    """
    Generate a self-contained ReportLab Drawing flowable containing a clean QR Code.
    Zero external dependencies — uses ReportLab's built-in QrCodeWidget.
    """
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    widget = QrCodeWidget(str(data))
    bounds = widget.getBounds()
    raw_w = max(bounds[2] - bounds[0], 1)
    raw_h = max(bounds[3] - bounds[1], 1)
    drawing = Drawing(size, size, transform=[size / raw_w, 0, 0, size / raw_h, 0, 0])
    drawing.add(widget)
    return drawing


def letterhead(doc_width, title, subtitle=None, filter_text=None, branding=None, logo_size=46,
               generated_at=None, generated_by=None):
    """Build the standard report letterhead: logo + site name/title + generation
    metadata right-aligned — the same structure reused by every report."""
    branding = branding or get_branding()
    regular, bold = document_fonts()
    style_title = ParagraphStyle("DDTitle", fontName=bold, fontSize=13, leading=17,
                                 textColor=colors.HexColor(INK))
    style_sub = ParagraphStyle("DDSub", fontName=regular, fontSize=8.5, leading=12,
                               textColor=colors.HexColor(MUTED))
    style_meta = ParagraphStyle("DDMeta", fontName=regular, fontSize=8, leading=11,
                                textColor=colors.HexColor(MUTED), alignment=2)

    from django.utils import timezone
    now_str = (generated_at or timezone.localtime()).strftime("%d %b %Y, %H:%M")
    meta_html = f"<b>Generated:</b> {now_str}"
    if filter_text:
        meta_html += f"<br/><b>Filter:</b> {escape(str(filter_text))}"
    if generated_by:
        meta_html += f"<br/><b>Generated by:</b> {escape(str(generated_by))}"

    title_html = f"<b>{escape(branding['site_name'].upper())}</b><br/>{escape(title)}"
    contact = branding.get("contact_line", "")
    sub_text = subtitle or contact or ""
    left_para = Paragraph(title_html, style_title)
    sub_para = Paragraph(escape(str(sub_text)), style_sub) if sub_text else None
    left_block = [left_para, sub_para] if sub_para else [left_para]

    logo = None
    if branding.get("logo_path"):
        try:
            logo = RLImage(branding["logo_path"])
            ratio = min(logo_size / logo.imageWidth, logo_size / logo.imageHeight)
            logo.drawWidth, logo.drawHeight = logo.imageWidth * ratio, logo.imageHeight * ratio
        except Exception:
            logo = None

    meta_width = min(200, doc_width * .32)
    if logo:
        header_rows = [[logo, left_block, Paragraph(meta_html, style_meta)]]
        col_widths = [logo_size + 10, doc_width - logo_size - 10 - meta_width, meta_width]
    else:
        header_rows = [[left_block, Paragraph(meta_html, style_meta)]]
        col_widths = [doc_width - meta_width, meta_width]

    table = Table(header_rows, colWidths=col_widths)
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def styled_table(data, col_widths, header_bg=HEADER_FILL, zebra=ZEBRA, border=BORDER, repeat_header=True):
    """Consistent table chrome (header fill, zebra rows, grid) for every report."""
    t = Table(data, colWidths=col_widths, repeatRows=1 if repeat_header else 0)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(zebra)]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(border)),
    ]))
    return t


# ==============================================================================
# EXCEL HELPERS
# ==============================================================================
def excel_title_banner(ws, title, subtitle, span_cols, primary=PRIMARY, primary_dark=PRIMARY_DARK):
    """The two-row merged title/subtitle banner used at the top of every export sheet."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    last_col = get_column_letter(span_cols)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"  {title}"
    ws["A1"].font = Font(name="Quicksand", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=primary.lstrip("#"), end_color=primary.lstrip("#"), fill_type="solid")
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 36

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"  {subtitle}"
    ws["A2"].font = Font(name="Quicksand", size=10, italic=True, color="E0E7FF")
    ws["A2"].fill = PatternFill(start_color=primary.lstrip("#"), end_color=primary.lstrip("#"), fill_type="solid")
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8


def excel_header_row(ws, row_idx, headers, primary_dark=PRIMARY_DARK):
    """Style the column-header row (row 4 by convention, after the title banner)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.row_dimensions[row_idx].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_num)
        cell.value = header
        cell.font = Font(name="Quicksand", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=primary_dark.lstrip("#"), end_color=primary_dark.lstrip("#"),
                                fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autofit_columns(ws, skip_rows=(1, 2, 3), min_width=11, max_width=45):
    """Approximate autofit column widths from cell content length."""
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in skip_rows:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def freeze_after_header(ws, header_row=4):
    """Freeze panes so the header row(s) stay visible while scrolling."""
    ws.freeze_panes = f"A{header_row + 1}"


def finish_worksheet(ws, header_row=4, data_end=None, filters=True):
    """Apply print, navigation, visible gridlines and readable cell formats."""
    from openpyxl.styles import Alignment, PatternFill, Border, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import MergedCell
    import math

    end = data_end if data_end is not None else ws.max_row
    last_col = get_column_letter(ws.max_column)
    freeze_after_header(ws, header_row)
    if filters and ws.max_row >= header_row:
        ws.auto_filter.ref = f'A{header_row}:{last_col}{max(header_row, end)}'

    # Ensure gridlines are clearly visible for tabular readability
    try:
        ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass
    ws.sheet_view.showGridLines = True

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if ws.max_column > 6 else 'portrait'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f'1:{header_row}'
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{last_col}{ws.max_row}'
    ws.page_margins = PageMargins(left=.3, right=.3, top=.45, bottom=.45, header=.2, footer=.2)
    ws.oddFooter.left.text = get_branding()['site_name'] + ' · Generated Official Report'
    ws.oddFooter.right.text = 'Page &P of &N'
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8
    autofit_columns(ws, skip_rows=range(1, header_row), max_width=44)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row in ws.iter_rows():
        lines = 1
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            font = copy(cell.font)
            font.name = 'Quicksand'
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.vertical = 'center'
            alignment.wrap_text = True
            cell.alignment = alignment
            if cell.row >= header_row:
                if not cell.border or not cell.border.left or cell.border.left.style is None:
                    cell.border = thin_border
                width = max(ws.column_dimensions[cell.column_letter].width - 2, 8)
                lines = max(lines, sum(max(1, math.ceil(len(s) / width)) for s in str(cell.value or '').split('\n')))
            if cell.row == header_row:
                font.bold, font.color = True, 'FFFFFF'
                cell.font = font
                cell.fill = PatternFill('solid', fgColor=PRIMARY_DARK.lstrip('#'))
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = 'dd mmm yyyy'
            elif isinstance(cell.value, (float, Decimal)) and cell.number_format == 'General':
                cell.number_format = '#,##0.00'
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool) and cell.number_format == 'General':
                cell.number_format = '#,##0'
        if row[0].row >= header_row:
            ws.row_dimensions[row[0].row].height = max(24, lines * 14 + 6)


# ==============================================================================
# CSV HELPERS
# ==============================================================================
def sanitize_csv_value(val):
    """Sanitize CSV cell content to stop formula injection when opened in spreadsheets."""
    if val is None:
        return ""
    s = str(val)
    if s.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + s
    return s


def build_clean_csv(headers, rows):
    """Generate a clean CSV byte stream with UTF-8 BOM encoding and formula escaping."""
    import csv
    import io
    buf = io.StringIO()
    buf.write("\ufeff")  # UTF-8 BOM
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([sanitize_csv_value(c) for c in row])
    return buf.getvalue().encode("utf-8")


def make_qr_drawing(url: str, size: float = 50.0):
    """
    Generate an authentic, scannable vector QR Code drawing for ReportLab PDFs.
    Zero external dependencies, uses ReportLab's native QrCodeWidget.
    """
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    try:
        widget = QrCodeWidget(str(url))
        bounds = widget.getBounds()
        w = bounds[2] - bounds[0]
        h = bounds[3] - bounds[1]
        if w <= 0 or h <= 0:
            w, h = 100.0, 100.0
        d = Drawing(size, size, transform=[size / w, 0, 0, size / h, 0, 0])
        d.add(widget)
        return d
    except Exception:
        return None

