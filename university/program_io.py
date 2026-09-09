import csv
import io
import re

from django.db import transaction
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from university.document_design import (
    ReportDocTemplate, document_styles, document_fonts,
    PageNumberCanvas, letterhead, get_branding, finish_worksheet
)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable, Image as RLImage, Paragraph, SimpleDocTemplate,
    Table, TableStyle, Spacer, KeepTogether
)

from university.models import Program, Department, School, AuditLog
from university.audit_services import log_activity


PROGRAM_IMPORT_COLUMNS = [
    ("code", "Programme Code", True),
    ("name", "Programme Name", True),
    ("department_code", "Department Code", True),
    ("level", "Level (CERT/DIP/UG/PG/MS/PHD)", True),
    ("program_type", "Programme Type", False),
    ("award_title", "Official Award Title", False),
    ("study_mode", "Study Mode", False),
    ("duration_value", "Duration Value", False),
    ("duration_unit", "Duration Unit (Years/Semesters)", False),
    ("min_credits", "Minimum Credits", False),
    ("total_seats", "Total Seats", False),
    ("status", "Status (Active/Inactive)", False),
    ("description", "Description", False),
]

SAMPLE_PROGRAM_ROW = {
    "code": "BS-CS",
    "name": "Bachelor of Science in Computer Science",
    "department_code": "CSE",
    "level": "UG",
    "program_type": "Degree",
    "award_title": "Bachelor of Science in Computer Science",
    "study_mode": "Full-Time",
    "duration_value": 4,
    "duration_unit": "Years",
    "min_credits": 128,
    "total_seats": 100,
    "status": "Active",
    "description": "Comprehensive four-year computer science program focusing on software engineering, data structures, and AI.",
}


# ==============================================================================
# 1. EXPORT FUNCTIONS (CSV, EXCEL, PDF)
# ==============================================================================

def export_program_csv(queryset):
    """Generate CSV byte string with UTF-8 BOM for programmes."""
    buffer = io.StringIO()
    buffer.write("\ufeff")
    writer = csv.writer(buffer)

    headers = [
        "Programme Code", "Programme Name", "Department Code", "Department Name",
        "Faculty / School", "Level", "Programme Type", "Award Title",
        "Study Mode", "Duration", "Min Credits", "Max Credits",
        "Status", "Enrolled Students", "Courses Count"
    ]
    writer.writerow(headers)

    for p in queryset:
        dept_code = p.department.code if p.department else "—"
        dept_name = p.department.name if p.department else "—"
        school_name = p.faculty_name
        students_count = getattr(p, "student_count", p.students.count())
        courses_count = getattr(p, "course_count", p.courses.count())
        duration_str = f"{p.duration_value} {p.duration_unit}"

        writer.writerow([
            p.code,
            p.name,
            dept_code,
            dept_name,
            school_name,
            p.get_level_display(),
            p.get_program_type_display(),
            p.award_title or p.name,
            p.get_study_mode_display(),
            duration_str,
            p.min_credits,
            p.max_credits or "—",
            p.status,
            students_count,
            courses_count,
        ])

    return buffer.getvalue().encode("utf-8")


def export_program_excel(queryset, site_name="University Management System"):
    """Generate professionally styled Excel (.xlsx) file for Programmes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programmes"
    ws.views.sheetView[0].showGridLines = True

    primary_color = "6C5CE7"      # Royal Indigo
    header_fill_color = "4834D4"  # Deep Indigo
    zebra_color = "EEF2FF"        # Light Indigo tint
    border_color = "CBD5E1"

    font_title = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=10, italic=True, color="E0E7FF")
    font_header = Font(name="Arial", size=10.5, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=9.5, color="1E293B")
    font_bold_data = Font(name="Arial", size=9.5, bold=True, color="1E293B")

    thin_border = Border(
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )

    # 1. Title Banner
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"  {site_name.upper()} — PROGRAMME DIRECTORY"
    title_cell.font = font_title
    title_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 36

    # 2. Subtitle / Metadata Row
    ws.merge_cells("A2:M2")
    sub_cell = ws["A2"]
    sub_cell.value = f"  Generated: {timezone.now().strftime('%d %b %Y, %H:%M')} | Total Programmes: {queryset.count()}"
    sub_cell.font = font_sub
    sub_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    sub_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 8

    # 3. Column Headers
    headers = [
        "Code", "Programme Name", "Department", "Faculty / School",
        "Level", "Type", "Study Mode", "Duration", "Min Credits",
        "Total Seats", "Status", "Students", "Courses"
    ]

    ws.row_dimensions[4].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
        cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in [1, 5, 6, 8, 9, 10, 11, 12, 13] else "left")
        cell.border = thin_border

    # 4. Data Rows
    current_row = 5
    for idx, p in enumerate(queryset):
        fill_color = zebra_color if idx % 2 == 1 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        students_count = getattr(p, "student_count", p.students.count())
        courses_count = getattr(p, "course_count", p.courses.count())
        duration_str = f"{p.duration_value} {p.duration_unit}"

        row_values = [
            p.code,
            p.name,
            p.department.code if p.department else "—",
            p.faculty_name,
            p.get_level_display(),
            p.get_program_type_display(),
            p.get_study_mode_display(),
            duration_str,
            p.min_credits,
            p.total_seats,
            p.status,
            students_count,
            courses_count,
        ]

        ws.row_dimensions[current_row].height = 21
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_bold_data if col_idx in [1, 12] else font_data
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [1, 3, 5, 6, 8, 9, 10, 11, 12, 13]:
                cell.alignment = Alignment(vertical="center", horizontal="center")
            else:
                cell.alignment = Alignment(vertical="center", horizontal="left")

        current_row += 1

    # 5. Auto-fit column widths
    finish_worksheet(ws, header_row=4)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_program_pdf(queryset, title="Programmes Directory", subtitle=None):
    """Generate formal A4 Landscape PDF for Programmes using ReportLab."""
    buffer = io.BytesIO()
    pagesize = landscape(A4)
    page_width, page_height = pagesize
    margin = 36

    doc = ReportDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )

    styles = document_styles()
    col_header_style = ParagraphStyle(
        "ProgHeader",
        parent=styles["Normal"],
        fontName="Quicksand-Bold",
        fontSize=8.5,
        textColor=colors.white,
        alignment=0,
    )
    cell_style = ParagraphStyle(
        "ProgCell",
        parent=styles["Normal"],
        fontName="Quicksand",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold_style = ParagraphStyle(
        "ProgCellBold",
        parent=cell_style,
        fontName="Quicksand-Bold",
    )
    cell_center_style = ParagraphStyle(
        "ProgCellCenter",
        parent=cell_style,
        alignment=1,
    )

    story = []
    sub = subtitle or f"Generated on {timezone.now().strftime('%d %B %Y, %H:%M')} · Total Programmes: {queryset.count()}"
    story.append(letterhead(page_width - 2 * margin, title, subtitle=sub))
    story.append(Spacer(1, 10))

    headers = [
        Paragraph("Code", col_header_style),
        Paragraph("Programme Name", col_header_style),
        Paragraph("Department", col_header_style),
        Paragraph("Faculty / School", col_header_style),
        Paragraph("Level", col_header_style),
        Paragraph("Mode", col_header_style),
        Paragraph("Duration", col_header_style),
        Paragraph("Credits", col_header_style),
        Paragraph("Seats", col_header_style),
        Paragraph("Status", col_header_style),
        Paragraph("Students", col_header_style),
    ]

    avail_w = page_width - 2 * margin
    col_fractions = [0.08, 0.24, 0.14, 0.14, 0.06, 0.08, 0.07, 0.05, 0.05, 0.04, 0.05]
    col_w = [avail_w * f for f in col_fractions]

    table_data = [headers]

    for p in queryset:
        students_count = getattr(p, "student_count", p.students.count())
        table_data.append([
            Paragraph(p.code, cell_bold_style),
            Paragraph(p.name, cell_style),
            Paragraph(p.department.name if p.department else "—", cell_style),
            Paragraph(p.faculty_name, cell_style),
            Paragraph(p.get_level_display(), cell_center_style),
            Paragraph(p.get_study_mode_display(), cell_center_style),
            Paragraph(f"{p.duration_value} {p.duration_unit}", cell_center_style),
            Paragraph(str(p.min_credits), cell_center_style),
            Paragraph(str(p.total_seats), cell_center_style),
            Paragraph(p.status, cell_center_style),
            Paragraph(str(students_count), cell_center_style),
        ])

    t = Table(table_data, colWidths=col_w, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4834D4")),
        ("ALIGN", (0, 0), (-1, 0), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FE")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))

    story.append(t)
    doc.build(story, canvasmaker=PageNumberCanvas)
    return buffer.getvalue()


# ==============================================================================
# 2. TEMPLATE GENERATION (CSV & EXCEL)
# ==============================================================================

def generate_program_import_template(fmt="csv"):
    """Generate blank template with sample row for Programme import."""
    headers = [col[0] for col in PROGRAM_IMPORT_COLUMNS]

    if fmt == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programmes Import Template"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
        font_header = Font(name="Arial", size=10.5, bold=True, color="FFFFFF")
        font_sample = Font(name="Arial", size=9.5, italic=True, color="64748B")

        ws.row_dimensions[1].height = 26
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", horizontal="center")

        ws.row_dimensions[2].height = 20
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=SAMPLE_PROGRAM_ROW.get(h, ""))
            cell.font = font_sample
            cell.alignment = Alignment(vertical="center", horizontal="left")

        finish_worksheet(ws, header_row=1, filters=False)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # CSV format
    buffer = io.StringIO()
    buffer.write("\ufeff")
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerow([SAMPLE_PROGRAM_ROW.get(h, "") for h in headers])
    return buffer.getvalue().encode("utf-8")


# ==============================================================================
# 3. IMPORT PARSER & VALIDATOR
# ==============================================================================

def parse_and_validate_program_import_file(uploaded_file):
    """
    Parse uploaded CSV or Excel file, validate each row against database constraints,
    and return structured validation results for preview.
    """
    filename = uploaded_file.name.lower()
    raw_rows = []

    if filename.endswith(".csv"):
        content = uploaded_file.read().decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(content))
        for row in reader:
            raw_rows.append({k.strip(): (v.strip() if v else "") for k, v in row.items() if k})
    elif filename.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return {"error": "The uploaded spreadsheet is empty."}

        headers = [str(h).strip() for h in header_row if h is not None]
        for r in rows_iter:
            if not any(r):
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                val = r[idx] if idx < len(r) else ""
                row_dict[h] = str(val).strip() if val is not None else ""
            raw_rows.append(row_dict)
    else:
        return {"error": "Unsupported file format. Please upload a .csv or .xlsx file."}

    if not raw_rows:
        return {"error": "The uploaded file contains no data rows."}

    # Pre-fetch lookup caches (case-insensitive)
    existing_codes = {c.strip().upper() for c in Program.objects.values_list("code", flat=True) if c}
    departments_by_code = {d.code.strip().upper(): d for d in Department.objects.all() if d.code}

    valid_rows = []
    invalid_rows = []
    seen_in_batch = set()

    for row_num, row in enumerate(raw_rows, start=2):
        code = row.get("code", "").strip().upper()
        name = row.get("name", "").strip()
        dept_code = row.get("department_code", "").strip().upper()
        level = row.get("level", "UG").strip().upper()
        prog_type = row.get("program_type", "Degree").strip()
        award_title = row.get("award_title", "").strip() or name
        study_mode = row.get("study_mode", "Full-Time").strip()
        duration_val = row.get("duration_value", "4").strip()
        duration_unit = row.get("duration_unit", "Years").strip()
        min_credits = row.get("min_credits", "120").strip()
        total_seats = row.get("total_seats", "120").strip()
        status = row.get("status", "Active").strip()
        description = row.get("description", "").strip()

        row_errors = []

        # Required fields
        if not code:
            row_errors.append("Programme Code is required.")
        elif code in existing_codes:
            row_errors.append(f"Programme Code '{code}' already exists in the system.")
        elif code in seen_in_batch:
            row_errors.append(f"Duplicate Programme Code '{code}' in uploaded file.")
        else:
            seen_in_batch.add(code)

        if not name:
            row_errors.append("Programme Name is required.")

        # Department validation
        dept_obj = departments_by_code.get(dept_code)
        if not dept_code:
            row_errors.append("Department Code is required.")
        elif not dept_obj:
            row_errors.append(f"Department Code '{dept_code}' does not exist.")

        # Level validation
        valid_levels = [choice[0] for choice in Program.LEVELS]
        if level not in valid_levels:
            row_errors.append(f"Invalid level '{level}'. Allowed: {', '.join(valid_levels)}")

        # Numbers validation
        try:
            dur_int = int(duration_val)
            if dur_int <= 0:
                row_errors.append("Duration must be a positive integer.")
        except ValueError:
            dur_int = 4
            row_errors.append("Duration must be a valid integer.")

        try:
            credits_int = int(min_credits)
            if credits_int <= 0:
                row_errors.append("Min Credits must be a positive integer.")
        except ValueError:
            credits_int = 120
            row_errors.append("Min Credits must be a valid integer.")

        try:
            seats_int = int(total_seats)
        except ValueError:
            seats_int = 120

        # Status validation
        if status.lower() in ["active", "1", "true", "yes"]:
            status = "Active"
        elif status.lower() in ["inactive", "0", "false", "no"]:
            status = "Inactive"
        elif status.lower() in ["phased out", "phased_out", "archived"]:
            status = "Phased Out"
        else:
            status = "Active"

        prepared_row = {
            "row_num": row_num,
            "code": code,
            "name": name,
            "department_id": dept_obj.id if dept_obj else None,
            "department_code": dept_code,
            "department_name": dept_obj.name if dept_obj else "—",
            "level": level,
            "program_type": prog_type,
            "award_title": award_title,
            "study_mode": study_mode,
            "duration_value": dur_int,
            "duration_unit": duration_unit.capitalize() if duration_unit else "Years",
            "min_credits": credits_int,
            "total_seats": seats_int,
            "status": status,
            "description": description,
            "errors": row_errors,
        }

        if row_errors:
            invalid_rows.append(prepared_row)
        else:
            valid_rows.append(prepared_row)

    return {
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "total_rows": len(raw_rows),
        "valid_count": len(valid_rows),
        "invalid_count": len(invalid_rows),
    }


# ==============================================================================
# 4. COMMIT IMPORT
# ==============================================================================

def commit_program_import(valid_rows, user=None):
    """Atomically insert validated programme records into the database."""
    created_programs = []

    with transaction.atomic():
        for r in valid_rows:
            p = Program.objects.create(
                code=r["code"],
                name=r["name"],
                department_id=r["department_id"],
                level=r["level"],
                program_type=r["program_type"],
                award_title=r["award_title"],
                study_mode=r["study_mode"],
                duration_value=r["duration_value"],
                duration_unit=r["duration_unit"],
                min_credits=r["min_credits"],
                total_seats=r["total_seats"],
                status=r["status"],
                description=r["description"],
            )
            created_programs.append(p)

        if user:
            log_activity(
                user=user,
                action=AuditLog.Action.IMPORT,
                module=AuditLog.Module.PROGRAMMES,
                entity="Program",
                description=f"Bulk imported {len(created_programs)} academic programmes.",
                new_state={"imported_codes": [p.code for p in created_programs]},
            )

    return len(created_programs)
